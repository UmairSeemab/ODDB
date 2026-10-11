#!/usr/bin/env python3
from __future__ import annotations

import argparse
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pymysql


CATEGORY_MAP = {
    "autosomal dominant": "INH-01",
    "autosomal recessive": "INH-02",
    "isolated cases": "INH-03",
    "mitochondrial": "INH-04",
    "other": "INH-05",
    "multiple": "INH-05",
    "unknown": "INH-05",
    "x-linked dominant": "INH-06",
    "x-linked recessive": "INH-07",
}

CATEGORY_NAMES = {
    "INH-01": "Autosomal dominant",
    "INH-02": "Autosomal recessive",
    "INH-03": "Isolated cases",
    "INH-04": "Mitochondrial",
    "INH-05": "Other",
    "INH-06": "X-linked dominant",
    "INH-07": "X-linked recessive",
}


def norm_header(x):
    return re.sub(r"[^a-z0-9]+", "", str(x or "").strip().lower())


def norm_name(x):
    x = unicodedata.normalize("NFKD", str(x or ""))
    x = "".join(ch for ch in x if not unicodedata.combining(ch))
    x = x.upper().strip()
    x = x.replace("’", "'").replace("–", "-").replace("—", "-")
    x = re.sub(r"\s*[-_/]\s*", " ", x)
    x = re.sub(r"[^A-Z0-9]+", " ", x)
    return re.sub(r"\s+", " ", x).strip()


def norm_mode(x):
    x = str(x or "").strip().lower()
    x = x.replace("_", " ").replace("–", "-").replace("—", "-")
    x = re.sub(r"\s+", " ", x)
    x = re.sub(r"^\d+\s*[.)-]?\s*", "", x)
    synonyms = {
        "autosomal dominant inheritance": "autosomal dominant",
        "autosomal recessive inheritance": "autosomal recessive",
        "isolated case": "isolated cases",
        "isolated": "isolated cases",
        "sporadic": "isolated cases",
        "maternal": "mitochondrial",
        "maternal inheritance": "mitochondrial",
        "x linked dominant": "x-linked dominant",
        "x linked recessive": "x-linked recessive",
        "others": "other",
    }
    return synonyms.get(x, x)


def norm_mondo(x):
    s = str(x or "").upper()
    m = re.search(r"MONDO[:_]?0*(\d+)", s)
    return f"MONDO:{int(m.group(1)):07d}" if m else None


def norm_orpha(x):
    s = str(x or "")
    m = re.search(r"(?:ORPHA(?:NET)?[:_ ]*)?(\d+)", s, re.I)
    return str(int(m.group(1))) if m else None


def norm_efo(x):
    s = str(x or "").upper()
    m = re.search(r"EFO[:_]?0*(\d+)", s)
    return f"EFO:{int(m.group(1)):07d}" if m else None


def find_column(headers, candidates):
    lookup = {norm_header(h): i for i, h in enumerate(headers) if h is not None}
    for c in candidates:
        k = norm_header(c)
        if k in lookup:
            return lookup[k]
    return None


def get_cell(row, idx):
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def detect_sheet(wb):
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration:
            continue
        if find_column(headers, ["MODE.OF.INHERITANCE", "MODE OF INHERITANCE", "Inheritance"]) is not None:
            return ws, headers
    raise SystemExit("No worksheet with an inheritance column was found.")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--report-dir", default="reports/web_taxonomy_inheritance_legacy_v0.3.8")
    ap.add_argument("--host", default=os.getenv("MYSQL_HOST", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(os.getenv("MYSQL_PORT", "3306")))
    ap.add_argument("--user", default=os.getenv("MYSQL_USER", "root"))
    ap.add_argument("--password", default=os.getenv("MYSQL_PASSWORD", ""))
    ap.add_argument("--database", default=os.getenv("MYSQL_DATABASE", "oddb_v1"))
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    if not xlsx.is_file():
        raise SystemExit(f"Workbook not found: {xlsx}")

    report_dir = Path(args.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws, headers = detect_sheet(wb)

    inh_i = find_column(headers, ["MODE.OF.INHERITANCE", "MODE OF INHERITANCE", "Inheritance"])
    did_i = find_column(headers, ["DiseaseID", "Disease ID", "Disease_ID", "ODDB_DiseaseID"])
    dname_i = find_column(headers, ["Ocular.disease", "Ocular disease", "Disease", "Disease name"])
    mondo_i = find_column(headers, ["EFO_Ids_Mondo", "MONDO", "MONDO_ID", "EFO MONDO"])
    orpha_i = find_column(headers, ["ORPHanet_ID", "Orphanet_ID", "ORPHA", "ORPHA_ID"])

    print("ODDB v1 LEGACY INHERITANCE XLSX IMPORTER v0.3.8")
    print("=" * 78)
    print(f"Workbook: {xlsx}")
    print(f"Worksheet: {ws.title}")
    print(f"Inheritance column: {headers[inh_i] if inh_i is not None else None}")
    print(f"Disease ID column: {headers[did_i] if did_i is not None else None}")
    print(f"Disease name column: {headers[dname_i] if dname_i is not None else None}")
    print(f"MONDO/EFO column: {headers[mondo_i] if mondo_i is not None else None}")
    print(f"Orphanet column: {headers[orpha_i] if orpha_i is not None else None}")

    conn = pymysql.connect(
        host=args.host, port=args.port, user=args.user, password=args.password,
        database=args.database, charset="utf8mb4", autocommit=False,
        cursorclass=pymysql.cursors.DictCursor,
    )

    try:
        with conn.cursor() as cur:
            # Discover disease schema.
            cur.execute("SHOW COLUMNS FROM disease")
            disease_cols = [r["Field"] for r in cur.fetchall()]
            disease_colset = set(disease_cols)

            select_cols = ["disease_id", "disease_name"]
            for c in ["mondo_id", "orphanet_id", "efo_id"]:
                if c in disease_colset:
                    select_cols.append(c)

            cur.execute("SELECT " + ",".join(f"`{c}`" for c in select_cols) + " FROM disease")
            diseases = cur.fetchall()

            by_id = {str(r["disease_id"]).strip(): r["disease_id"] for r in diseases}
            by_name = defaultdict(set)
            by_mondo = defaultdict(set)
            by_orpha = defaultdict(set)
            by_efo = defaultdict(set)

            for r in diseases:
                did = r["disease_id"]
                by_name[norm_name(r["disease_name"])].add(did)

                if "mondo_id" in r and r.get("mondo_id"):
                    n = norm_mondo(r["mondo_id"])
                    if n:
                        by_mondo[n].add(did)

                if "orphanet_id" in r and r.get("orphanet_id"):
                    n = norm_orpha(r["orphanet_id"])
                    if n:
                        by_orpha[n].add(did)

                if "efo_id" in r and r.get("efo_id"):
                    n = norm_efo(r["efo_id"])
                    if n:
                        by_efo[n].add(did)

            # Discover disease_alias schema.
            cur.execute("SHOW COLUMNS FROM disease_alias")
            alias_cols = [r["Field"] for r in cur.fetchall()]
            alias_lookup = {norm_header(c): c for c in alias_cols}

            alias_text_col = None
            for candidate in ["alias_name", "alias", "disease_alias", "alias_text",
                              "disease_name_alias", "old_disease_name", "historical_name"]:
                if norm_header(candidate) in alias_lookup:
                    alias_text_col = alias_lookup[norm_header(candidate)]
                    break

            alias_source_id_col = None
            for candidate in ["source_id", "alias_id", "old_disease_id",
                              "previous_disease_id", "legacy_disease_id"]:
                if norm_header(candidate) in alias_lookup:
                    alias_source_id_col = alias_lookup[norm_header(candidate)]
                    break

            old_to_current = {}
            if alias_text_col:
                cur.execute(
                    f"SELECT disease_id,`{alias_text_col}` alias_text "
                    f"FROM disease_alias WHERE `{alias_text_col}` IS NOT NULL "
                    f"AND `{alias_text_col}`<>''"
                )
                for r in cur.fetchall():
                    by_name[norm_name(r["alias_text"])].add(r["disease_id"])

            if alias_source_id_col:
                cur.execute(
                    f"SELECT disease_id,`{alias_source_id_col}` legacy_id "
                    f"FROM disease_alias WHERE `{alias_source_id_col}` IS NOT NULL "
                    f"AND `{alias_source_id_col}`<>''"
                )
                for r in cur.fetchall():
                    old_to_current[str(r["legacy_id"]).strip()] = r["disease_id"]

            old_to_current.setdefault("ODDB-76", "ODDB-225")

            raw_mode_counts = Counter()
            mapping_method_counts = Counter()
            unmapped_modes = Counter()
            mapped = defaultdict(lambda: defaultdict(set))
            unresolved = []
            ambiguous = []
            seen = set()
            duplicate_source_rows = 0

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                raw_mode = get_cell(row, inh_i)
                if not raw_mode:
                    continue

                mode = norm_mode(raw_mode)
                raw_mode_counts[raw_mode] += 1
                cid = CATEGORY_MAP.get(mode)
                if not cid:
                    unmapped_modes[raw_mode] += 1
                    continue

                did_raw = get_cell(row, did_i)
                dname_raw = get_cell(row, dname_i)
                mondo_raw = get_cell(row, mondo_i)
                orpha_raw = get_cell(row, orpha_i)

                source_key = (did_raw, dname_raw, raw_mode, mondo_raw, orpha_raw)
                if source_key in seen:
                    duplicate_source_rows += 1
                seen.add(source_key)

                candidates = []
                methods = []

                # 1. Current stable ODDB DiseaseID.
                if did_raw in by_id:
                    candidates.append(by_id[did_raw])
                    methods.append("current DiseaseID")

                # 2. Retired/source ODDB ID.
                if did_raw in old_to_current:
                    candidates.append(old_to_current[did_raw])
                    methods.append("legacy/source DiseaseID")

                # 3. Exact MONDO identifier.
                mondo = norm_mondo(mondo_raw)
                if mondo and mondo in by_mondo:
                    candidates.extend(by_mondo[mondo])
                    methods.append("MONDO")

                # EFO can coexist in EFO_Ids_Mondo.
                efo = norm_efo(mondo_raw)
                if efo and efo in by_efo:
                    candidates.extend(by_efo[efo])
                    methods.append("EFO")

                # 4. Exact Orphanet identifier.
                orpha = norm_orpha(orpha_raw)
                if orpha and orpha in by_orpha:
                    candidates.extend(by_orpha[orpha])
                    methods.append("Orphanet")

                # 5. Exact normalized current/alias name.
                if dname_raw:
                    candidates.extend(by_name.get(norm_name(dname_raw), set()))
                    if by_name.get(norm_name(dname_raw)):
                        methods.append("exact name/alias")

                unique = sorted(set(candidates))

                if len(unique) == 1:
                    resolved = unique[0]
                    method = "+".join(sorted(set(methods)))
                    mapping_method_counts[method] += 1
                    mapped[resolved][cid].add(
                        f"{method}; legacy DiseaseID={did_raw}; "
                        f"raw inheritance={raw_mode}"
                    )
                elif len(unique) > 1:
                    ambiguous.append((
                        row_num, did_raw, dname_raw, mondo_raw, orpha_raw,
                        raw_mode, "|".join(unique), "|".join(sorted(set(methods)))
                    ))
                else:
                    unresolved.append((
                        row_num, did_raw, dname_raw, mondo_raw, orpha_raw, raw_mode
                    ))

            membership_count = sum(len(cats) for cats in mapped.values())

            print(f"\nSource rows with inheritance: {sum(raw_mode_counts.values()):,}")
            print(f"Unique mapped current diseases: {len(mapped):,}")
            print(f"Mapped disease/category memberships: {membership_count:,}")
            print(f"Duplicate source rows: {duplicate_source_rows:,}")
            print(f"Unresolved source rows: {len(unresolved):,}")
            print(f"Ambiguous source rows: {len(ambiguous):,}")

            print("\nMapping methods:")
            for k, n in mapping_method_counts.most_common():
                print(f"{n:6d}  {k}")

            print("\nRaw inheritance values:")
            for k, n in raw_mode_counts.most_common():
                print(f"{n:6d}  {k}")

            if unmapped_modes:
                print("\nUNMAPPED inheritance labels:")
                for k, n in unmapped_modes.most_common():
                    print(f"{n:6d}  {k}")

            counts = Counter()
            for did, cats in mapped.items():
                for cid in cats:
                    counts[cid] += 1

            print("\nPreview category counts:")
            for cid in CATEGORY_NAMES:
                print(f"{CATEGORY_NAMES[cid]:28s} {counts[cid]:4d}")

            # Reports.
            with (report_dir / "mapped_memberships.tsv").open("w", encoding="utf-8") as f:
                f.write("DiseaseID\tCategory_ID\tCategory_Name\tMembership_Basis\n")
                for did in sorted(mapped):
                    for cid in sorted(mapped[did]):
                        f.write(
                            f"{did}\t{cid}\t{CATEGORY_NAMES[cid]}\t"
                            + " || ".join(sorted(mapped[did][cid]))
                            + "\n"
                        )

            with (report_dir / "unresolved_rows.tsv").open("w", encoding="utf-8") as f:
                f.write("Row\tDiseaseID\tDisease_Name\tMONDO_EFO\tOrphanet\tInheritance\n")
                for r in unresolved:
                    f.write("\t".join(map(str, r)) + "\n")

            with (report_dir / "ambiguous_rows.tsv").open("w", encoding="utf-8") as f:
                f.write(
                    "Row\tDiseaseID\tDisease_Name\tMONDO_EFO\tOrphanet\tInheritance\t"
                    "Candidate_DiseaseIDs\tMethods\n"
                )
                for r in ambiguous:
                    f.write("\t".join(map(str, r)) + "\n")

            if not args.apply:
                print("\nPREVIEW ONLY — database was not changed.")
                print(f"Reports: {report_dir}")
                return

            cur.execute("""
                DELETE m
                FROM disease_category_membership m
                JOIN disease_category c ON c.category_id=m.category_id
                WHERE c.category_system='inheritance'
            """)

            inserted = 0
            for did in sorted(mapped):
                for cid in sorted(mapped[did]):
                    basis = " || ".join(sorted(mapped[did][cid]))
                    cur.execute("""
                        INSERT INTO disease_category_membership
                          (disease_id,category_id,membership_basis,source,curation_status)
                        VALUES (%s,%s,%s,'ODDB legacy Inheritance_Based_file.xlsx','legacy_curated')
                    """, (did, cid, basis[:255]))
                    inserted += 1

            conn.commit()

            print(f"\nAPPLIED: {inserted:,} inheritance memberships written to MySQL.")
            print("Raw values Multiple/Unknown were represented under the UI category Other")
            print("and preserved in membership provenance.")
            print(f"Reports: {report_dir}")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
