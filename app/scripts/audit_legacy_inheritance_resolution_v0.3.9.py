#!/usr/bin/env python3
from __future__ import annotations

import argparse
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pymysql


def norm_header(x):
    return re.sub(r"[^a-z0-9]+", "", str(x or "").strip().lower())


def norm_name(x):
    x = unicodedata.normalize("NFKD", str(x or ""))
    x = "".join(ch for ch in x if not unicodedata.combining(ch))
    x = x.upper().strip()
    x = x.replace("’", "'").replace("–", "-").replace("—", "-")
    x = re.sub(r"\s*[-_/]\s*", " ", x)
    x = re.sub(r"[^A-Z0-9]+", " ", x)
    return re.sub(r"\s+", " ", x).strip()


def normalize_oddb_id(x):
    """
    Safe syntactic normalizations only:
      ODDB_25 -> ODDB-25
      ODDB 25 -> ODDB-25
      25 / 25.0 -> ODDB-25
    """
    s = str(x or "").strip().upper()
    if not s:
        return None

    m = re.fullmatch(r"ODDB[\s_-]*(\d+)(?:\.0+)?", s)
    if m:
        return f"ODDB-{int(m.group(1))}"

    m = re.fullmatch(r"(\d+)(?:\.0+)?", s)
    if m:
        return f"ODDB-{int(m.group(1))}"

    return s


def extract_mondo_values(x):
    s = str(x or "").upper()
    vals = set()
    for m in re.finditer(r"MONDO[:_ ]*0*(\d+)", s):
        vals.add(f"MONDO:{int(m.group(1)):07d}")
    return vals


def extract_efo_values(x):
    s = str(x or "").upper()
    vals = set()
    for m in re.finditer(r"EFO[:_ ]*0*(\d+)", s):
        vals.add(f"EFO:{int(m.group(1)):07d}")
    return vals


def extract_orpha_values(x):
    s = str(x or "")
    vals = set()

    # Explicit ORPHA/Orphanet identifiers.
    for m in re.finditer(r"(?:ORPHA(?:NET)?)[\s:_-]*0*(\d+)", s, re.I):
        vals.add(str(int(m.group(1))))

    # If field appears to contain only numeric IDs separated by punctuation.
    if not vals and re.fullmatch(r"[\d\s,;|/.()-]+", s.strip()):
        for m in re.finditer(r"\d+", s):
            vals.add(str(int(m.group(0))))
    return vals


def normalize_db_mondo(x):
    vals = extract_mondo_values(x)
    return vals


def normalize_db_orpha(x):
    return extract_orpha_values(x)


def find_column(headers, candidates):
    lookup = {norm_header(h): i for i, h in enumerate(headers) if h is not None}
    for c in candidates:
        if norm_header(c) in lookup:
            return lookup[norm_header(c)]
    return None


def cell(row, idx):
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--report-dir", default="reports/web_taxonomy_inheritance_resolution_v0.3.9")
    ap.add_argument("--host", default=os.getenv("MYSQL_HOST", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(os.getenv("MYSQL_PORT", "3306")))
    ap.add_argument("--user", default=os.getenv("MYSQL_USER", "root"))
    ap.add_argument("--password", default=os.getenv("MYSQL_PASSWORD", ""))
    ap.add_argument("--database", default=os.getenv("MYSQL_DATABASE", "oddb_v1"))
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    report_dir = Path(args.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    ws = None
    headers = None
    for candidate in wb.worksheets:
        it = candidate.iter_rows(values_only=True)
        try:
            h = list(next(it))
        except StopIteration:
            continue
        if find_column(h, ["MODE.OF.INHERITANCE", "Inheritance"]) is not None:
            ws, headers = candidate, h
            break

    if ws is None:
        raise SystemExit("No inheritance worksheet found.")

    did_i = find_column(headers, ["DiseaseID", "Disease ID", "Disease_ID"])
    name_i = find_column(headers, ["Ocular.disease", "Ocular disease", "Disease", "Disease name"])
    inh_i = find_column(headers, ["MODE.OF.INHERITANCE", "Inheritance"])
    mondo_i = find_column(headers, ["EFO_Ids_Mondo", "MONDO", "MONDO_ID"])
    orpha_i = find_column(headers, ["ORPHanet_ID", "Orphanet_ID", "ORPHA"])

    conn = pymysql.connect(
        host=args.host, port=args.port, user=args.user, password=args.password,
        database=args.database, charset="utf8mb4", autocommit=False,
        cursorclass=pymysql.cursors.DictCursor
    )

    try:
        with conn.cursor() as cur:
            cur.execute("SHOW COLUMNS FROM disease")
            disease_cols = {r["Field"] for r in cur.fetchall()}

            select_cols = ["disease_id", "disease_name"]
            for c in ["mondo_id", "orphanet_id", "efo_id"]:
                if c in disease_cols:
                    select_cols.append(c)

            cur.execute("SELECT " + ",".join(f"`{c}`" for c in select_cols) + " FROM disease")
            db_rows = cur.fetchall()

            current_ids = {r["disease_id"] for r in db_rows}
            by_name = defaultdict(set)
            by_mondo = defaultdict(set)
            by_orpha = defaultdict(set)
            by_efo = defaultdict(set)

            for r in db_rows:
                did = r["disease_id"]
                by_name[norm_name(r["disease_name"])].add(did)

                for v in normalize_db_mondo(r.get("mondo_id")):
                    by_mondo[v].add(did)
                for v in normalize_db_orpha(r.get("orphanet_id")):
                    by_orpha[v].add(did)
                for v in extract_efo_values(r.get("efo_id")):
                    by_efo[v].add(did)

            # Alias support.
            cur.execute("SHOW COLUMNS FROM disease_alias")
            alias_cols = [r["Field"] for r in cur.fetchall()]
            alias_lookup = {norm_header(c): c for c in alias_cols}

            alias_col = alias_lookup.get(norm_header("alias"))
            source_id_col = alias_lookup.get(norm_header("source_id"))

            old_to_current = {}
            if alias_col:
                cur.execute(
                    f"SELECT disease_id,`{alias_col}` alias_text "
                    f"FROM disease_alias WHERE `{alias_col}` IS NOT NULL AND `{alias_col}`<>''"
                )
                for r in cur.fetchall():
                    by_name[norm_name(r["alias_text"])].add(r["disease_id"])

            if source_id_col:
                cur.execute(
                    f"SELECT disease_id,`{source_id_col}` legacy_id "
                    f"FROM disease_alias WHERE `{source_id_col}` IS NOT NULL AND `{source_id_col}`<>''"
                )
                for r in cur.fetchall():
                    old_to_current[normalize_oddb_id(r["legacy_id"])] = r["disease_id"]

            old_to_current.setdefault("ODDB-76", "ODDB-225")

            source_rows = []
            raw_id_patterns = Counter()

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                inh = cell(row, inh_i)
                if not inh:
                    continue

                did_raw = cell(row, did_i)
                name_raw = cell(row, name_i)
                mondo_raw = cell(row, mondo_i)
                orpha_raw = cell(row, orpha_i)

                if not did_raw:
                    raw_id_patterns["<blank>"] += 1
                elif re.fullmatch(r"ODDB-\d+", did_raw, re.I):
                    raw_id_patterns["ODDB-N"] += 1
                elif re.fullmatch(r"\d+(?:\.0+)?", did_raw):
                    raw_id_patterns["numeric"] += 1
                else:
                    raw_id_patterns["other"] += 1

                source_rows.append({
                    "row": row_num,
                    "did_raw": did_raw,
                    "did_norm": normalize_oddb_id(did_raw),
                    "name": name_raw,
                    "inh": inh,
                    "mondo_raw": mondo_raw,
                    "orpha_raw": orpha_raw,
                })

            # Source-level uniqueness diagnostics.
            unique_ids = {r["did_raw"] for r in source_rows if r["did_raw"]}
            unique_norm_ids = {r["did_norm"] for r in source_rows if r["did_norm"]}
            unique_names = {norm_name(r["name"]) for r in source_rows if r["name"]}
            unique_source_diseases = {
                (r["did_norm"] or "", norm_name(r["name"]))
                for r in source_rows
                if r["did_norm"] or r["name"]
            }

            resolution_counts = Counter()
            resolved_current_diseases = set()
            unresolved_unique = {}
            conflicts = []

            for r in source_rows:
                evidence = defaultdict(set)

                # Stable ID, including safe syntactic normalization.
                if r["did_norm"] in current_ids:
                    evidence["DiseaseID"].add(r["did_norm"])
                if r["did_norm"] in old_to_current:
                    evidence["legacy/source DiseaseID"].add(old_to_current[r["did_norm"]])

                # Exact identifiers: allow multiple IDs in one cell.
                for m in extract_mondo_values(r["mondo_raw"]):
                    for did in by_mondo.get(m, set()):
                        evidence["MONDO"].add(did)

                for e in extract_efo_values(r["mondo_raw"]):
                    for did in by_efo.get(e, set()):
                        evidence["EFO"].add(did)

                for o in extract_orpha_values(r["orpha_raw"]):
                    for did in by_orpha.get(o, set()):
                        evidence["Orphanet"].add(did)

                if r["name"]:
                    for did in by_name.get(norm_name(r["name"]), set()):
                        evidence["exact name/alias"].add(did)

                all_hits = set()
                for vals in evidence.values():
                    all_hits.update(vals)

                if len(all_hits) == 1:
                    did = next(iter(all_hits))
                    resolved_current_diseases.add(did)
                    method = "+".join(sorted(k for k, v in evidence.items() if did in v))
                    resolution_counts[method] += 1
                elif len(all_hits) > 1:
                    conflicts.append((r, evidence, sorted(all_hits)))
                else:
                    key = (r["did_norm"] or "", norm_name(r["name"]))
                    unresolved_unique.setdefault(key, r)

            print("ODDB LEGACY INHERITANCE RESOLUTION AUDIT v0.3.9")
            print("=" * 80)
            print(f"Workbook: {xlsx}")
            print(f"Worksheet: {ws.title}")
            print(f"Source rows with inheritance: {len(source_rows):,}")
            print(f"Unique raw DiseaseID values: {len(unique_ids):,}")
            print(f"Unique normalized DiseaseID values: {len(unique_norm_ids):,}")
            print(f"Unique disease-name values: {len(unique_names):,}")
            print(f"Unique source disease keys: {len(unique_source_diseases):,}")
            print(f"Current ODDB diseases: {len(current_ids):,}")
            print(f"Resolved current ODDB diseases: {len(resolved_current_diseases):,}")
            print(f"Unresolved unique source disease keys: {len(unresolved_unique):,}")
            print(f"Conflicting source rows: {len(conflicts):,}")

            print("\nLegacy DiseaseID formatting:")
            for k, n in raw_id_patterns.most_common():
                print(f"{k:20s} {n:6d}")

            print("\nResolution methods by source row:")
            for k, n in resolution_counts.most_common():
                print(f"{n:6d}  {k}")

            print("\nFirst 40 unresolved unique disease keys:")
            for i, (_, r) in enumerate(sorted(unresolved_unique.items())[:40], start=1):
                print(
                    f"{i:2d}. row={r['row']}  DiseaseID={r['did_raw']!r}  "
                    f"normalized={r['did_norm']!r}  disease={r['name']!r}  "
                    f"inheritance={r['inh']!r}  MONDO/EFO={r['mondo_raw']!r}  "
                    f"Orphanet={r['orpha_raw']!r}"
                )

            unresolved_path = report_dir / "unresolved_unique_diseases.tsv"
            with unresolved_path.open("w", encoding="utf-8") as f:
                f.write("Row\tDiseaseID_Raw\tDiseaseID_Normalized\tDisease_Name\tInheritance\tMONDO_EFO\tOrphanet\n")
                for _, r in sorted(unresolved_unique.items()):
                    f.write(
                        f"{r['row']}\t{r['did_raw']}\t{r['did_norm'] or ''}\t"
                        f"{r['name']}\t{r['inh']}\t{r['mondo_raw']}\t{r['orpha_raw']}\n"
                    )

            conflict_path = report_dir / "conflicting_rows.tsv"
            with conflict_path.open("w", encoding="utf-8") as f:
                f.write("Row\tDiseaseID\tDisease_Name\tInheritance\tCandidate_DiseaseIDs\tEvidence\n")
                for r, evidence, hits in conflicts:
                    ev = "; ".join(
                        f"{k}={','.join(sorted(v))}" for k, v in evidence.items() if v
                    )
                    f.write(
                        f"{r['row']}\t{r['did_raw']}\t{r['name']}\t{r['inh']}\t"
                        f"{'|'.join(hits)}\t{ev}\n"
                    )

            print("\nReports:")
            print(unresolved_path)
            print(conflict_path)
            print("\nREAD-ONLY audit: MySQL was not modified.")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
