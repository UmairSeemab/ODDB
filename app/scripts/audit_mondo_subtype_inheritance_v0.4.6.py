#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import os
import re
from collections import Counter, defaultdict, deque
from pathlib import Path

import openpyxl
import pymysql


CATEGORY_NAMES = {
    "INH-01": "Autosomal dominant",
    "INH-02": "Autosomal recessive",
    "INH-04": "Mitochondrial",
    "INH-05": "Other",
    "INH-06": "X-linked dominant",
    "INH-07": "X-linked recessive",
}


def map_inheritance(raw):
    s = str(raw or "").strip().lower()
    s = s.replace("_", " ").replace("–", "-").replace("—", "-")
    s = re.sub(r"\s+", " ", s)

    if "autosomal dominant" in s:
        return "INH-01"
    if "autosomal recessive" in s:
        return "INH-02"
    if "mitochondrial" in s:
        return "INH-04"
    if "x-linked dominant" in s or "x linked dominant" in s:
        return "INH-06"
    if "x-linked recessive" in s or "x linked recessive" in s:
        return "INH-07"
    if s in {"multiple", "unknown", "other", "others"}:
        return "INH-05"
    return None


def norm_header(x):
    return re.sub(r"[^a-z0-9]+", "", str(x or "").strip().lower())


def find_col(headers, candidates):
    lookup = {norm_header(h): i for i, h in enumerate(headers) if h is not None}
    for c in candidates:
        k = norm_header(c)
        if k in lookup:
            return lookup[k]
    return None


def cell(row, idx):
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def extract_mondo_values(x):
    s = str(x or "").upper()
    vals = set()
    for m in re.finditer(r"MONDO[:_ ]*0*(\d+)", s):
        vals.add(f"MONDO:{int(m.group(1)):07d}")
    return vals


def parse_mondo_obo(path):
    """
    Minimal OBO parser for [Term] id, name and is_a hierarchy.
    We deliberately use only asserted is_a parentage.
    """
    parents = defaultdict(set)
    names = {}

    current_id = None
    current_name = None
    in_term = False

    def flush():
        nonlocal current_id, current_name
        if current_id and current_name:
            names[current_id] = current_name
        current_id = None
        current_name = None

    with open(path, "r", encoding="utf-8", errors="replace") as f:
        for raw in f:
            line = raw.rstrip("\n")

            if line == "[Term]":
                if in_term:
                    flush()
                in_term = True
                continue

            if line.startswith("[") and line.endswith("]") and line != "[Term]":
                if in_term:
                    flush()
                in_term = False
                continue

            if not in_term:
                continue

            if not line.strip():
                flush()
                in_term = False
                continue

            if line.startswith("id: "):
                current_id = line[4:].strip()
            elif line.startswith("name: "):
                current_name = line[6:].strip()
            elif line.startswith("is_a: ") and current_id:
                parent = line[6:].split(" ! ", 1)[0].strip()
                if parent.startswith("MONDO:"):
                    parents[current_id].add(parent)

    if in_term:
        flush()

    return parents, names


def ancestor_distances(node, parents):
    """
    Return shortest asserted is_a distance from node to every ancestor.
    distance 1 = direct parent.
    """
    dist = {}
    q = deque([(node, 0)])

    while q:
        cur, d = q.popleft()
        for p in parents.get(cur, ()):
            nd = d + 1
            if p not in dist or nd < dist[p]:
                dist[p] = nd
                q.append((p, nd))
    return dist


def shortest_path_to_ancestor(node, ancestor, parents):
    q = deque([(node, [node])])
    seen = {node}

    while q:
        cur, path = q.popleft()
        for p in parents.get(cur, ()):
            if p == ancestor:
                return path + [p]
            if p not in seen:
                seen.add(p)
                q.append((p, path + [p]))
    return []


def read_missing(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return {
            r["DiseaseID"]
            for r in csv.DictReader(f, delimiter="\t")
            if r.get("DiseaseID")
        }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--legacy-xlsx", required=True)
    ap.add_argument("--mondo", required=True)
    ap.add_argument(
        "--missing",
        default="reports/web_taxonomy_inheritance_hpo_v0.4.4/"
                "still_missing_after_hpo.tsv",
    )
    ap.add_argument(
        "--report-dir",
        default="reports/web_taxonomy_inheritance_mondo_subtypes_v0.4.6",
    )
    ap.add_argument("--host", default=os.getenv("MYSQL_HOST", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(os.getenv("MYSQL_PORT", "3306")))
    ap.add_argument("--user", default=os.getenv("MYSQL_USER", "root"))
    ap.add_argument("--password", default=os.getenv("MYSQL_PASSWORD", ""))
    ap.add_argument("--database", default=os.getenv("MYSQL_DATABASE", "oddb_v1"))
    args = ap.parse_args()

    legacy_xlsx = Path(args.legacy_xlsx)
    mondo_path = Path(args.mondo)

    if not legacy_xlsx.is_file():
        raise SystemExit(f"Legacy workbook not found: {legacy_xlsx}")
    if not mondo_path.is_file():
        raise SystemExit(f"MONDO file not found: {mondo_path}")

    missing_ids = read_missing(args.missing)
    report_dir = Path(args.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    print("Parsing MONDO asserted hierarchy...")
    parents, mondo_names = parse_mondo_obo(mondo_path)
    print(f"MONDO terms with names: {len(mondo_names):,}")
    print(f"MONDO terms with asserted parent(s): {len(parents):,}")

    conn = pymysql.connect(
        host=args.host,
        port=args.port,
        user=args.user,
        password=args.password,
        database=args.database,
        charset="utf8mb4",
        autocommit=False,
        cursorclass=pymysql.cursors.DictCursor,
    )

    try:
        with conn.cursor() as cur:
            cur.execute("SHOW COLUMNS FROM disease")
            dcols = {r["Field"] for r in cur.fetchall()}
            if "mondo_id" not in dcols:
                raise SystemExit("Current disease table has no mondo_id column.")

            cur.execute("""
                SELECT disease_id, disease_name, mondo_id
                FROM disease
            """)
            current = {
                r["disease_id"]: r
                for r in cur.fetchall()
                if r["disease_id"] in missing_ids
            }

        current_by_mondo = defaultdict(set)
        for did, d in current.items():
            for m in extract_mondo_values(d.get("mondo_id")):
                current_by_mondo[m].add(did)

        # Read legacy inheritance workbook.
        wb = openpyxl.load_workbook(
            legacy_xlsx, read_only=True, data_only=True
        )
        ws = wb[wb.sheetnames[0]]
        headers = list(next(ws.iter_rows(values_only=True)))

        did_i = find_col(headers, ["DiseaseID"])
        name_i = find_col(headers, ["Ocular.disease", "Ocular disease", "Disease"])
        inh_i = find_col(headers, ["MODE.OF.INHERITANCE", "Inheritance"])
        mondo_i = find_col(headers, ["EFO_Ids_Mondo", "MONDO", "MONDO_ID"])

        legacy_rows = []
        for row_no, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2
        ):
            raw_inh = cell(row, inh_i)
            cid = map_inheritance(raw_inh)
            if not cid:
                continue

            mondo_vals = extract_mondo_values(cell(row, mondo_i))
            if not mondo_vals:
                continue

            for lm in mondo_vals:
                legacy_rows.append({
                    "row": row_no,
                    "legacy_disease_id": cell(row, did_i),
                    "legacy_disease_name": cell(row, name_i),
                    "legacy_mondo": lm,
                    "raw_inheritance": raw_inh,
                    "category_id": cid,
                })

        candidates = []
        summary_support = defaultdict(list)
        distance_counts = Counter()
        legacy_with_parent_match = set()

        # Cache ancestry because legacy MONDOs repeat.
        ancestry_cache = {}

        for lr in legacy_rows:
            lm = lr["legacy_mondo"]

            if lm not in ancestry_cache:
                ancestry_cache[lm] = ancestor_distances(lm, parents)

            ancestors = ancestry_cache[lm]

            for parent_mondo, distance in ancestors.items():
                if parent_mondo not in current_by_mondo:
                    continue

                # This is deliberately subtype/ancestor support, never identity.
                for current_did in current_by_mondo[parent_mondo]:
                    cd = current[current_did]
                    path = shortest_path_to_ancestor(lm, parent_mondo, parents)

                    rec = {
                        "Current_DiseaseID": current_did,
                        "Current_Disease_Name": cd["disease_name"],
                        "Current_MONDO": parent_mondo,
                        "Category_ID": lr["category_id"],
                        "Category_Name": CATEGORY_NAMES[lr["category_id"]],
                        "Legacy_DiseaseID": lr["legacy_disease_id"],
                        "Legacy_Disease_Name": lr["legacy_disease_name"],
                        "Legacy_MONDO": lm,
                        "Legacy_Inheritance_Raw": lr["raw_inheritance"],
                        "MONDO_Distance": distance,
                        "MONDO_Path": " > ".join(path),
                        "Legacy_Source_Row": lr["row"],
                        "Evidence_Type": "subtype-supported",
                    }
                    candidates.append(rec)
                    summary_support[
                        (current_did, lr["category_id"])
                    ].append(rec)
                    distance_counts[distance] += 1
                    legacy_with_parent_match.add(
                        (lr["legacy_disease_id"], lm)
                    )

        newly_coverable = {
            did for (did, _cid) in summary_support
        }
        still_missing = set(current) - newly_coverable

        print("\nODDB v1 MONDO SUBTYPE-INHERITANCE AUDIT v0.4.6")
        print("=" * 86)
        print(f"Current diseases missing direct inheritance: {len(current):,}")
        print(f"Legacy inheritance rows with MONDO IDs: {len(legacy_rows):,}")
        print(f"Subtype-support candidate evidence rows: {len(candidates):,}")
        print(f"Current diseases with >=1 subtype-supported category: "
              f"{len(newly_coverable):,}")
        print(f"Still missing after hypothetical subtype support: "
              f"{len(still_missing):,}")
        print(f"Candidate current disease/category memberships: "
              f"{len(summary_support):,}")

        print("\nCandidate evidence by MONDO ancestor distance:")
        for d, n in sorted(distance_counts.items()):
            print(f"distance {d:2d}: {n:6d}")

        cat_counts = Counter()
        for (_did, cid) in summary_support:
            cat_counts[cid] += 1

        print("\nCandidate memberships by inheritance category:")
        for cid in CATEGORY_NAMES:
            print(f"{CATEGORY_NAMES[cid]:28s} {cat_counts[cid]:4d}")

        print("\nFirst 30 newly coverable diseases:")
        for did in sorted(newly_coverable)[:30]:
            cats = sorted({
                CATEGORY_NAMES[cid]
                for (d, cid) in summary_support
                if d == did
            })
            min_dist = min(
                rec["MONDO_Distance"]
                for (d, _cid), recs in summary_support.items()
                if d == did
                for rec in recs
            )
            print(
                f"{did:10s}  {current[did]['disease_name'][:46]:46s} "
                f"min_distance={min_dist:2d}  "
                f"categories={' | '.join(cats)}"
            )

        fieldnames = [
            "Current_DiseaseID", "Current_Disease_Name", "Current_MONDO",
            "Category_ID", "Category_Name",
            "Legacy_DiseaseID", "Legacy_Disease_Name", "Legacy_MONDO",
            "Legacy_Inheritance_Raw",
            "MONDO_Distance", "MONDO_Path",
            "Legacy_Source_Row", "Evidence_Type",
        ]

        with (report_dir / "subtype_supported_candidate_evidence.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.DictWriter(f, fieldnames=fieldnames, delimiter="\t")
            w.writeheader()
            w.writerows(sorted(
                candidates,
                key=lambda r: (
                    r["Current_DiseaseID"],
                    r["Category_ID"],
                    r["MONDO_Distance"],
                    r["Legacy_DiseaseID"],
                )
            ))

        with (report_dir / "candidate_membership_summary.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow([
                "Current_DiseaseID", "Current_Disease_Name",
                "Category_ID", "Category_Name",
                "Supporting_Legacy_Subtypes",
                "Minimum_MONDO_Distance",
                "Legacy_Subtype_Examples",
            ])
            for (did, cid), recs in sorted(summary_support.items()):
                examples = sorted({
                    f"{r['Legacy_DiseaseID']}:{r['Legacy_Disease_Name']}"
                    for r in recs
                })
                w.writerow([
                    did,
                    current[did]["disease_name"],
                    cid,
                    CATEGORY_NAMES[cid],
                    len({
                        (r["Legacy_DiseaseID"], r["Legacy_MONDO"])
                        for r in recs
                    }),
                    min(r["MONDO_Distance"] for r in recs),
                    " | ".join(examples[:10]),
                ])

        with (report_dir / "still_missing_after_subtype_candidates.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow(["DiseaseID", "Disease_Name", "MONDO_ID"])
            for did in sorted(still_missing):
                w.writerow([
                    did,
                    current[did]["disease_name"],
                    current[did].get("mondo_id", ""),
                ])

        print("\nReports:")
        print(report_dir / "subtype_supported_candidate_evidence.tsv")
        print(report_dir / "candidate_membership_summary.tsv")
        print(report_dir / "still_missing_after_subtype_candidates.tsv")
        print("\nREAD-ONLY: MySQL was not modified.")
        print(
            "Important: these are subtype-supported browsing candidates, "
            "not direct disease-identity or direct disease-level inheritance assertions."
        )

    finally:
        conn.close()


if __name__ == "__main__":
    main()
