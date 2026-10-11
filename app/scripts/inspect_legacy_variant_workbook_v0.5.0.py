#!/usr/bin/env python3
from __future__ import annotations

import argparse
import os
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


def norm_header(x):
    return re.sub(r"[^a-z0-9]+", "", str(x or "").strip().lower())


def looks_like_url(x):
    s = str(x or "").strip().lower()
    return s.startswith("http://") or s.startswith("https://")


def main():
    ap = argparse.ArgumentParser(
        description="Read-only inspection of legacy ODDB Variant_Disease_final_file.xlsx"
    )
    ap.add_argument("--xlsx", required=True)
    ap.add_argument(
        "--report-dir",
        default="reports/legacy_variant_workbook_audit_v0.5.0",
    )
    ap.add_argument("--sample-rows", type=int, default=20)
    ap.add_argument("--top-values", type=int, default=40)
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    if not xlsx.is_file():
        raise SystemExit(f"Workbook not found: {xlsx}")

    report_dir = Path(args.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    print("ODDB LEGACY VARIANT WORKBOOK AUDIT v0.5.0")
    print("=" * 88)
    print(f"Workbook: {xlsx}")
    print(f"File size: {xlsx.stat().st_size / (1024*1024):.2f} MB")

    # Normal mode is intentional so true Excel hyperlinks can be inspected.
    wb = openpyxl.load_workbook(xlsx, read_only=False, data_only=False)

    print(f"Worksheets: {wb.sheetnames}")

    workbook_summary = []
    all_hyperlinks = []

    for ws in wb.worksheets:
        print("\n" + "=" * 88)
        print(f"Worksheet: {ws.title}")
        print(f"Dimensions: rows={ws.max_row:,}, columns={ws.max_column:,}")

        headers = [
            str(ws.cell(1, c).value).strip() if ws.cell(1, c).value is not None else ""
            for c in range(1, ws.max_column + 1)
        ]

        print("Headers:")
        for i, h in enumerate(headers, start=1):
            print(f"  {i:2d}. {h}")

        header_lookup = {norm_header(h): i for i, h in enumerate(headers, start=1) if h}

        candidate_groups = {
            "disease": [
                "DiseaseID", "Disease ID", "Disease", "Ocular.disease",
                "Ocular disease", "Disease name"
            ],
            "gene": [
                "Gene", "GeneSymbol", "Gene Symbol", "HGNC", "HGNC_ID",
                "KNOWN GENES OR CHROMOSOMAL ABNORMALITY INVOLVED"
            ],
            "variant_id": [
                "VariationID", "Variation ID", "VariantID", "Variant ID",
                "ClinVar Variation ID", "rsID", "dbSNP", "dbSNP ID",
                "Variant"
            ],
            "consequence": [
                "Consequence", "Variant consequence", "Variant_consequence",
                "Molecular consequence", "Molecular Consequence",
                "Annotation", "Effect", "Variant type", "Variant Type"
            ],
            "chromosome": [
                "Chromosome", "Chr", "Chrom"
            ],
            "position": [
                "Position", "Start", "Location", "Genomic position",
                "Genomic Position"
            ],
            "clinical_significance": [
                "Clinical significance", "Clinical Significance",
                "Significance", "Pathogenicity"
            ],
            "inheritance": [
                "Inheritance", "MODE.OF.INHERITANCE", "Mode of inheritance"
            ],
        }

        detected = {}
        for group, candidates in candidate_groups.items():
            cols = []
            for c in candidates:
                k = norm_header(c)
                if k in header_lookup:
                    cols.append((header_lookup[k], headers[header_lookup[k]-1]))
            detected[group] = cols

        print("\nDetected candidate columns:")
        for group, cols in detected.items():
            if cols:
                print(f"  {group:22s}: " + ", ".join(f"{idx}:{name}" for idx, name in cols))
            else:
                print(f"  {group:22s}: <none>")

        # Basic column profiling.
        value_counts = {
            c: Counter()
            for c in range(1, ws.max_column + 1)
        }
        nonempty_counts = Counter()
        url_text_counts = Counter()
        hyperlink_counts = Counter()
        unique_disease_ids = set()
        unique_gene_values = set()
        unique_variant_values = set()

        disease_cols = {idx for idx, _ in detected["disease"]}
        gene_cols = {idx for idx, _ in detected["gene"]}
        variant_cols = {idx for idx, _ in detected["variant_id"]}
        consequence_cols = {idx for idx, _ in detected["consequence"]}

        sample_rows = []

        for r in range(2, ws.max_row + 1):
            row_sample = {}

            for c in range(1, ws.max_column + 1):
                cell = ws.cell(r, c)
                val = cell.value
                if val is None or str(val).strip() == "":
                    continue

                sval = str(val).strip()
                nonempty_counts[c] += 1

                # Limit retained categorical frequencies to reasonable string lengths.
                if len(sval) <= 250:
                    value_counts[c][sval] += 1

                if looks_like_url(sval):
                    url_text_counts[c] += 1

                if cell.hyperlink and cell.hyperlink.target:
                    hyperlink_counts[c] += 1
                    all_hyperlinks.append({
                        "sheet": ws.title,
                        "row": r,
                        "column": c,
                        "header": headers[c-1],
                        "display": sval,
                        "target": cell.hyperlink.target,
                    })

                if c in disease_cols:
                    unique_disease_ids.add(sval)
                if c in gene_cols:
                    unique_gene_values.add(sval)
                if c in variant_cols:
                    unique_variant_values.add(sval)

                if r <= args.sample_rows + 1:
                    row_sample[headers[c-1] or f"Column_{c}"] = sval

            if row_sample and r <= args.sample_rows + 1:
                sample_rows.append((r, row_sample))

        print("\nColumn summary:")
        for c, h in enumerate(headers, start=1):
            print(
                f"  {c:2d}. {h[:42]:42s} "
                f"nonempty={nonempty_counts[c]:7d} "
                f"url_text={url_text_counts[c]:6d} "
                f"excel_hyperlinks={hyperlink_counts[c]:6d} "
                f"unique_seen={len(value_counts[c]):7d}"
            )

        print("\nPotential categorical columns / top values:")
        for c, h in enumerate(headers, start=1):
            nuniq = len(value_counts[c])
            nnon = nonempty_counts[c]

            # Show likely categorical columns: explicit consequence candidates,
            # or low-cardinality populated columns.
            if c in consequence_cols or (nnon > 0 and nuniq <= 100):
                print(f"\n[{c}] {h}  unique={nuniq:,}")
                for val, n in value_counts[c].most_common(args.top_values):
                    print(f"  {n:7d}  {val}")

        print("\nDetected identifier counts:")
        print(f"  Unique disease-like values: {len(unique_disease_ids):,}")
        print(f"  Unique gene-like values:    {len(unique_gene_values):,}")
        print(f"  Unique variant-like values: {len(unique_variant_values):,}")
        print(f"  Excel hyperlinks:           {sum(hyperlink_counts.values()):,}")
        print(f"  URL-text cells:             {sum(url_text_counts.values()):,}")

        print("\nSample rows:")
        for row_num, sample in sample_rows:
            print(f"\n  Row {row_num}:")
            for k, v in sample.items():
                print(f"    {k}: {v}")

        # Reports per sheet.
        safe_sheet = re.sub(r"[^A-Za-z0-9_.-]+", "_", ws.title)

        with (report_dir / f"{safe_sheet}_column_summary.tsv").open(
            "w", encoding="utf-8"
        ) as f:
            f.write(
                "Column\tHeader\tNonempty\tUnique_Values_Seen\t"
                "URL_Text_Cells\tExcel_Hyperlinks\n"
            )
            for c, h in enumerate(headers, start=1):
                f.write(
                    f"{c}\t{h}\t{nonempty_counts[c]}\t"
                    f"{len(value_counts[c])}\t{url_text_counts[c]}\t"
                    f"{hyperlink_counts[c]}\n"
                )

        with (report_dir / f"{safe_sheet}_top_values.tsv").open(
            "w", encoding="utf-8"
        ) as f:
            f.write("Column\tHeader\tValue\tCount\n")
            for c, h in enumerate(headers, start=1):
                nuniq = len(value_counts[c])
                nnon = nonempty_counts[c]
                if c in consequence_cols or (nnon > 0 and nuniq <= 100):
                    for val, n in value_counts[c].most_common(args.top_values):
                        f.write(f"{c}\t{h}\t{val}\t{n}\n")

        workbook_summary.append({
            "sheet": ws.title,
            "rows": ws.max_row,
            "columns": ws.max_column,
            "headers": headers,
            "unique_disease_like": len(unique_disease_ids),
            "unique_gene_like": len(unique_gene_values),
            "unique_variant_like": len(unique_variant_values),
            "hyperlinks": sum(hyperlink_counts.values()),
            "url_text": sum(url_text_counts.values()),
        })

    # Global hyperlink report.
    with (report_dir / "all_excel_hyperlinks.tsv").open("w", encoding="utf-8") as f:
        f.write("Sheet\tRow\tColumn\tHeader\tDisplay_Text\tTarget\n")
        for x in all_hyperlinks:
            f.write(
                f"{x['sheet']}\t{x['row']}\t{x['column']}\t{x['header']}\t"
                f"{x['display']}\t{x['target']}\n"
            )

    with (report_dir / "workbook_summary.tsv").open("w", encoding="utf-8") as f:
        f.write(
            "Sheet\tRows\tColumns\tUnique_Disease_Like\tUnique_Gene_Like\t"
            "Unique_Variant_Like\tExcel_Hyperlinks\tURL_Text_Cells\n"
        )
        for x in workbook_summary:
            f.write(
                f"{x['sheet']}\t{x['rows']}\t{x['columns']}\t"
                f"{x['unique_disease_like']}\t{x['unique_gene_like']}\t"
                f"{x['unique_variant_like']}\t{x['hyperlinks']}\t{x['url_text']}\n"
            )

    print("\n" + "=" * 88)
    print("Reports:")
    print(report_dir / "workbook_summary.tsv")
    print(report_dir / "all_excel_hyperlinks.tsv")
    for x in workbook_summary:
        safe_sheet = re.sub(r"[^A-Za-z0-9_.-]+", "_", x["sheet"])
        print(report_dir / f"{safe_sheet}_column_summary.tsv")
        print(report_dir / f"{safe_sheet}_top_values.tsv")

    print("\nREAD-ONLY: workbook and MySQL were not modified.")


if __name__ == "__main__":
    main()
