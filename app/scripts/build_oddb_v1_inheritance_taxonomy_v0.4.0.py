#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pymysql


CATEGORY_MAP = {
    "autosomal dominant": "INH-01",
    "autosomal recessive": "INH-02",
    "isolated cases": "INH-03",
    "mitochondrial": "INH-04",
    "other": "INH-05",
    "multiple": "INH-05",
    "unknown": "INH-05",
    "x-linked dominant": "INH-06",
    "x-linked recessive": "INH-07",
}

CATEGORY_NAMES = {
    "INH-01": "Autosomal dominant",
    "INH-02": "Autosomal recessive",
    "INH-03": "Isolated cases",
    "INH-04": "Mitochondrial",
    "INH-05": "Other",
    "INH-06": "X-linked dominant",
    "INH-07": "X-linked recessive",
}


def norm_header(x):
    return re.sub(r"[^a-z0-9]+", "", str(x or "").strip().lower())


def norm_name(x):
    x = unicodedata.normalize("NFKD", str(x or ""))
    x = "".join(ch for ch in x if not unicodedata.combining(ch))
    x = x.upper().strip()
    x = x.replace("’", "'").replace("–", "-").replace("—", "-")
    x = re.sub(r"\s*[-_/]\s*", " ", x)
    x = re.sub(r"[^A-Z0-9]+", " ", x)
    return re.sub(r"\s+", " ", x).strip()


def normalize_oddb_id(x):
    s = str(x or "").strip().upper()
    if not s:
        return None
    m = re.fullmatch(r"ODDB[\s_-]*(\d+)(?:\.0+)?", s)
    if m:
        return f"ODDB-{int(m.group(1))}"
    m = re.fullmatch(r"(\d+)(?:\.0+)?", s)
    if m:
        return f"ODDB-{int(m.group(1))}"
    return s


def extract_mondo_values(x):
    s = str(x or "").upper()
    vals = set()
    for m in re.finditer(r"MONDO[:_ ]*0*(\d+)", s):
        vals.add(f"MONDO:{int(m.group(1)):07d}")
    return vals


def extract_orpha_values(x):
    s = str(x or "")
    vals = set()
    for m in re.finditer(r"(?:ORPHA(?:NET)?)[\s:_-]*0*(\d+)", s, re.I):
        vals.add(str(int(m.group(1))))
    if not vals and re.fullmatch(r"[\d\s,;|/.()-]+", s.strip()):
        for m in re.finditer(r"\d+", s):
            vals.add(str(int(m.group(0))))
    return vals


def map_inheritance(raw):
    s = str(raw or "").strip().lower()
    if not s:
        return None
    s = s.replace("_", " ").replace("–", "-").replace("—", "-")
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"^\d+\s*[.)-]?\s*", "", s)

    # Explicit first.
    if "autosomal dominant" in s or re.fullmatch(r"ad", s):
        return "INH-01"
    if "autosomal recessive" in s or re.fullmatch(r"ar", s):
        return "INH-02"
    if "isolated case" in s or s == "isolated" or "sporadic" in s:
        return "INH-03"
    if "mitochondrial" in s or "maternal inheritance" in s or s == "maternal":
        return "INH-04"
    if "x-linked dominant" in s or "x linked dominant" in s or re.fullmatch(r"xld|xd", s):
        return "INH-06"
    if "x-linked recessive" in s or "x linked recessive" in s or re.fullmatch(r"xlr|xr", s):
        return "INH-07"

    # Generic/aggregate values belong to Other; do not infer a specific mode.
    if s in {
        "multiple", "unknown", "other", "others", "complex", "multifactorial",
        "x-linked", "x linked", "digenic", "oligogenic"
    }:
        return "INH-05"

    return None


def find_col(headers, candidates):
    lookup = {norm_header(h): i for i, h in enumerate(headers) if h is not None}
    for c in candidates:
        k = norm_header(c)
        if k in lookup:
            return lookup[k]
    return None


def row_cell(row, idx):
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def load_tsv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def find_field(fieldnames, candidates):
    lookup = {norm_header(f): f for f in fieldnames if f}
    for c in candidates:
        k = norm_header(c)
        if k in lookup:
            return lookup[k]
    return None


def main():
    ap = argparse.ArgumentParser(
        description="Build current ODDB v1 inheritance taxonomy for the 162 canonical diseases."
    )
    ap.add_argument("--legacy-xlsx", required=True)
    ap.add_argument("--retnet", required=True)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--report-dir", default="reports/web_taxonomy_inheritance_current_v0.4.0")
    ap.add_argument("--host", default=os.getenv("MYSQL_HOST", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(os.getenv("MYSQL_PORT", "3306")))
    ap.add_argument("--user", default=os.getenv("MYSQL_USER", "root"))
    ap.add_argument("--password", default=os.getenv("MYSQL_PASSWORD", ""))
    ap.add_argument("--database", default=os.getenv("MYSQL_DATABASE", "oddb_v1"))
    args = ap.parse_args()

    legacy_xlsx = Path(args.legacy_xlsx)
    retnet_tsv = Path(args.retnet)
    report_dir = Path(args.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    if not legacy_xlsx.is_file():
        raise SystemExit(f"Legacy inheritance workbook not found: {legacy_xlsx}")
    if not retnet_tsv.is_file():
        raise SystemExit(f"RetNet candidate file not found: {retnet_tsv}")

    conn = pymysql.connect(
        host=args.host, port=args.port, user=args.user, password=args.password,
        database=args.database, charset="utf8mb4", autocommit=False,
        cursorclass=pymysql.cursors.DictCursor,
    )

    try:
        with conn.cursor() as cur:
            # -----------------------------
            # Current disease universe
            # -----------------------------
            cur.execute("SHOW COLUMNS FROM disease")
            disease_cols = [r["Field"] for r in cur.fetchall()]
            dset = set(disease_cols)

            select_cols = ["disease_id", "disease_name"]
            for c in ["mondo_id", "orphanet_id"]:
                if c in dset:
                    select_cols.append(c)

            cur.execute("SELECT " + ",".join(f"`{c}`" for c in select_cols) + " FROM disease")
            diseases = cur.fetchall()

            current_ids = {r["disease_id"] for r in diseases}
            by_name = defaultdict(set)
            by_mondo = defaultdict(set)
            by_orpha = defaultdict(set)

            for r in diseases:
                did = r["disease_id"]
                by_name[norm_name(r["disease_name"])].add(did)
                for m in extract_mondo_values(r.get("mondo_id")):
                    by_mondo[m].add(did)
                for o in extract_orpha_values(r.get("orphanet_id")):
                    by_orpha[o].add(did)

            # -----------------------------
            # Disease aliases / retired IDs
            # -----------------------------
            cur.execute("SHOW COLUMNS FROM disease_alias")
            alias_cols = [r["Field"] for r in cur.fetchall()]
            alias_lookup = {norm_header(c): c for c in alias_cols}

            alias_col = alias_lookup.get("alias")
            source_id_col = alias_lookup.get("sourceid")
            old_to_current = {}

            if alias_col:
                cur.execute(
                    f"SELECT disease_id,`{alias_col}` AS alias_text "
                    f"FROM disease_alias WHERE `{alias_col}` IS NOT NULL AND `{alias_col}`<>''"
                )
                for r in cur.fetchall():
                    by_name[norm_name(r["alias_text"])].add(r["disease_id"])

            if source_id_col:
                cur.execute(
                    f"SELECT disease_id,`{source_id_col}` AS legacy_id "
                    f"FROM disease_alias WHERE `{source_id_col}` IS NOT NULL AND `{source_id_col}`<>''"
                )
                for r in cur.fetchall():
                    lid = normalize_oddb_id(r["legacy_id"])
                    if lid:
                        old_to_current[lid] = r["disease_id"]

            # Known retired alias from the frozen core curation.
            old_to_current.setdefault("ODDB-76", "ODDB-225")

            # Each membership keeps all independent evidence records.
            evidence = defaultdict(lambda: defaultdict(list))

            # -----------------------------
            # Source 1: legacy XLSX
            # exact mapping only
            # -----------------------------
            wb = openpyxl.load_workbook(legacy_xlsx, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            header_row = next(ws.iter_rows(values_only=True))
            headers = list(header_row)

            did_i = find_col(headers, ["DiseaseID"])
            name_i = find_col(headers, ["Ocular.disease", "Ocular disease", "Disease"])
            inh_i = find_col(headers, ["MODE.OF.INHERITANCE", "Inheritance"])
            mondo_i = find_col(headers, ["EFO_Ids_Mondo", "MONDO", "MONDO_ID"])
            orpha_i = find_col(headers, ["ORPHanet_ID", "Orphanet_ID", "ORPHA"])

            legacy_stats = Counter()
            legacy_unresolved = []
            legacy_ambiguous = []

            for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                raw_inh = row_cell(row, inh_i)
                if not raw_inh:
                    continue

                cid = map_inheritance(raw_inh)
                if not cid:
                    legacy_stats["unmapped_inheritance"] += 1
                    continue

                did_raw = normalize_oddb_id(row_cell(row, did_i))
                name_raw = row_cell(row, name_i)
                mondo_raw = row_cell(row, mondo_i)
                orpha_raw = row_cell(row, orpha_i)

                hits = set()
                methods = []

                if did_raw in current_ids:
                    hits.add(did_raw)
                    methods.append("DiseaseID")
                if did_raw in old_to_current:
                    hits.add(old_to_current[did_raw])
                    methods.append("retired/source DiseaseID")

                for m in extract_mondo_values(mondo_raw):
                    vals = by_mondo.get(m, set())
                    if vals:
                        hits.update(vals)
                        methods.append("MONDO")

                for o in extract_orpha_values(orpha_raw):
                    vals = by_orpha.get(o, set())
                    if vals:
                        hits.update(vals)
                        methods.append("Orphanet")

                if name_raw:
                    vals = by_name.get(norm_name(name_raw), set())
                    if vals:
                        hits.update(vals)
                        methods.append("exact name/alias")

                if len(hits) == 1:
                    did = next(iter(hits))
                    evidence[did][cid].append({
                        "source": "ODDB legacy Inheritance_Based_file.xlsx",
                        "raw_value": raw_inh,
                        "basis": "+".join(sorted(set(methods))),
                        "legacy_disease_id": did_raw or "",
                        "legacy_disease_name": name_raw,
                        "source_row": row_no,
                    })
                    legacy_stats["mapped_rows"] += 1
                elif len(hits) > 1:
                    legacy_ambiguous.append((row_no, did_raw, name_raw, raw_inh, "|".join(sorted(hits))))
                    legacy_stats["ambiguous_rows"] += 1
                else:
                    legacy_unresolved.append((row_no, did_raw, name_raw, raw_inh, mondo_raw, orpha_raw))
                    legacy_stats["unresolved_rows"] += 1

            # -----------------------------
            # Source 2: current DGA
            # explicit inheritance only
            # -----------------------------
            cur.execute("SHOW COLUMNS FROM disease_gene_association")
            dga_cols = [r["Field"] for r in cur.fetchall()]
            dga_lookup = {norm_header(c): c for c in dga_cols}

            dga_disease_col = dga_lookup.get("diseaseid") or dga_lookup.get("disease_id")
            dga_inh_col = dga_lookup.get("inheritance")

            dga_stats = Counter()
            if dga_disease_col and dga_inh_col:
                cur.execute(
                    f"SELECT `{dga_disease_col}` AS disease_id, "
                    f"`{dga_inh_col}` AS inheritance "
                    f"FROM disease_gene_association "
                    f"WHERE `{dga_inh_col}` IS NOT NULL AND `{dga_inh_col}`<>''"
                )
                for r in cur.fetchall():
                    did = r["disease_id"]
                    raw_inh = r["inheritance"]
                    cid = map_inheritance(raw_inh)
                    if did in current_ids and cid:
                        evidence[did][cid].append({
                            "source": "ODDB v1 disease_gene_association",
                            "raw_value": raw_inh,
                            "basis": "explicit DGA inheritance",
                            "legacy_disease_id": "",
                            "legacy_disease_name": "",
                            "source_row": "",
                        })
                        dga_stats["mapped_rows"] += 1
                    elif raw_inh:
                        dga_stats["unmapped_rows"] += 1
            else:
                dga_stats["inheritance_column_not_found"] += 1

            # -----------------------------
            # Source 3: RetNet exact mappings only
            # -----------------------------
            retnet_rows = load_tsv(retnet_tsv)
            if not retnet_rows:
                raise SystemExit("RetNet file is empty.")

            fields = retnet_rows[0].keys()
            r_did = find_field(fields, ["ODDB_DiseaseID", "DiseaseID"])
            r_inh = find_field(fields, ["Inheritance"])
            r_status = find_field(fields, ["Disease_Match_Status", "Disease Match Status"])

            retnet_stats = Counter()

            for r in retnet_rows:
                if r_status and str(r.get(r_status, "")).strip().lower() != "exact":
                    retnet_stats["skipped_nonexact"] += 1
                    continue

                did = normalize_oddb_id(r.get(r_did, "")) if r_did else None
                raw_inh = r.get(r_inh, "") if r_inh else ""
                cid = map_inheritance(raw_inh)

                if did in current_ids and cid:
                    evidence[did][cid].append({
                        "source": "RetNet exact mapping",
                        "raw_value": raw_inh,
                        "basis": "Disease_Match_Status=exact",
                        "legacy_disease_id": "",
                        "legacy_disease_name": "",
                        "source_row": "",
                    })
                    retnet_stats["mapped_rows"] += 1
                elif did in current_ids and raw_inh:
                    retnet_stats["unmapped_inheritance"] += 1
                elif did:
                    retnet_stats["unknown_current_id"] += 1

            # -----------------------------
            # Summaries
            # -----------------------------
            category_counts = Counter()
            covered = set()
            multi = {}
            for did, cats in evidence.items():
                if cats:
                    covered.add(did)
                if len(cats) > 1:
                    multi[did] = sorted(cats)
                for cid in cats:
                    category_counts[cid] += 1

            print("ODDB v1 CURRENT INHERITANCE TAXONOMY BUILDER v0.4.0")
            print("=" * 84)
            print(f"Current canonical diseases: {len(current_ids):,}")
            print(f"Diseases with >=1 supported inheritance category: {len(covered):,}")
            print(f"Diseases without supported inheritance category: {len(current_ids - covered):,}")
            print(f"Diseases with >1 inheritance category: {len(multi):,}")

            print("\nInheritance category counts:")
            for cid in CATEGORY_NAMES:
                print(f"{CATEGORY_NAMES[cid]:28s} {category_counts[cid]:4d}")

            print("\nSource diagnostics:")
            print("Legacy XLSX:", dict(legacy_stats))
            print("Current DGA:", dict(dga_stats))
            print("RetNet exact:", dict(retnet_stats))

            # -----------------------------
            # Reports
            # -----------------------------
            with (report_dir / "current_inheritance_memberships.tsv").open("w", encoding="utf-8") as f:
                f.write("DiseaseID\tCategory_ID\tCategory_Name\tEvidence_Count\tSources\tRaw_Values\tEvidence_JSON\n")
                for did in sorted(evidence):
                    for cid in sorted(evidence[did]):
                        ev = evidence[did][cid]
                        sources = sorted({x["source"] for x in ev})
                        raws = sorted({str(x["raw_value"]) for x in ev if str(x["raw_value"]).strip()})
                        f.write(
                            f"{did}\t{cid}\t{CATEGORY_NAMES[cid]}\t{len(ev)}\t"
                            f"{' | '.join(sources)}\t{' | '.join(raws)}\t"
                            f"{json.dumps(ev, ensure_ascii=False)}\n"
                        )

            name_by_id = {r["disease_id"]: r["disease_name"] for r in diseases}

            with (report_dir / "inheritance_missing_diseases.tsv").open("w", encoding="utf-8") as f:
                f.write("DiseaseID\tDisease_Name\n")
                for did in sorted(current_ids - covered):
                    f.write(f"{did}\t{name_by_id.get(did,'')}\n")

            with (report_dir / "inheritance_multi_category_diseases.tsv").open("w", encoding="utf-8") as f:
                f.write("DiseaseID\tDisease_Name\tCategories\n")
                for did, cats in sorted(multi.items()):
                    f.write(
                        f"{did}\t{name_by_id.get(did,'')}\t"
                        + " | ".join(CATEGORY_NAMES[c] for c in cats)
                        + "\n"
                    )

            with (report_dir / "legacy_unresolved_rows.tsv").open("w", encoding="utf-8") as f:
                f.write("Row\tLegacy_DiseaseID\tLegacy_Disease_Name\tInheritance\tMONDO_EFO\tOrphanet\n")
                for r in legacy_unresolved:
                    f.write("\t".join(map(str, r)) + "\n")

            with (report_dir / "legacy_ambiguous_rows.tsv").open("w", encoding="utf-8") as f:
                f.write("Row\tLegacy_DiseaseID\tLegacy_Disease_Name\tInheritance\tCandidate_Current_DiseaseIDs\n")
                for r in legacy_ambiguous:
                    f.write("\t".join(map(str, r)) + "\n")

            print("\nReports:")
            for p in [
                "current_inheritance_memberships.tsv",
                "inheritance_missing_diseases.tsv",
                "inheritance_multi_category_diseases.tsv",
                "legacy_unresolved_rows.tsv",
                "legacy_ambiguous_rows.tsv",
            ]:
                print(report_dir / p)

            if not args.apply:
                print("\nPREVIEW ONLY — MySQL was not modified.")
                return

            # Replace only current inheritance web taxonomy memberships.
            cur.execute("""
                DELETE m
                FROM disease_category_membership m
                JOIN disease_category c ON c.category_id=m.category_id
                WHERE c.category_system='inheritance'
            """)

            inserted = 0
            for did in sorted(evidence):
                for cid in sorted(evidence[did]):
                    ev = evidence[did][cid]
                    sources = sorted({x["source"] for x in ev})
                    raws = sorted({str(x["raw_value"]) for x in ev if str(x["raw_value"]).strip()})

                    membership_basis = (
                        "sources=" + " | ".join(sources)
                        + "; raw=" + " | ".join(raws)
                    )

                    cur.execute("""
                        INSERT INTO disease_category_membership
                            (disease_id,category_id,membership_basis,source,curation_status)
                        VALUES (%s,%s,%s,%s,%s)
                    """, (
                        did,
                        cid,
                        membership_basis[:255],
                        "ODDB v1 inheritance harmonization v0.4.0",
                        "supported",
                    ))
                    inserted += 1

            conn.commit()

            print(f"\nAPPLIED: {inserted:,} current inheritance memberships written to MySQL.")
            print("Only the inheritance web-taxonomy memberships were replaced.")
            print("Canonical Disease/Gene/Variant/Drug/clinical tables were not changed.")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
