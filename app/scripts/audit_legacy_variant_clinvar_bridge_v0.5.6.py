#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pymysql


EXPECTED_VARIANT_TYPES = [
    "missense variant",
    "inframe deletion",
    "frameshift variant",
    "intron variant",
    "regulatory region variant",
    "intergenic variant",
    "splice region variant",
    "splice donor variant",
    "non coding transcript exon variant",
    "3 prime UTR variant",
    "5 prime UTR variant",
    "stop gained",
    "synonymous variant",
    "TF binding site variant",
    "splice acceptor variant",
    "downstream gene variant",
    "stop lost",
    "upstream gene variant",
    "inframe insertion",
    "protein altering variant",
    "Multiple",
]


def norm_header(x):
    return re.sub(r"[^a-z0-9]+", "", str(x or "").strip().lower())


def norm_name(x):
    x = unicodedata.normalize("NFKD", str(x or ""))
    x = "".join(ch for ch in x if not unicodedata.combining(ch))
    x = x.upper().strip()
    x = x.replace("’", "'").replace("–", "-").replace("—", "-")
    x = re.sub(r"\s*[-_/]\s*", " ", x)
    x = re.sub(r"[^A-Z0-9]+", " ", x)
    return re.sub(r"\s+", " ", x).strip()


def normalize_oddb_id(x):
    s = str(x or "").strip().upper()
    if not s:
        return None
    m = re.fullmatch(r"ODDB[\s_-]*(\d+)(?:\.0+)?", s)
    if m:
        return f"ODDB-{int(m.group(1))}"
    return s


def normalize_rsid(value, allow_bare_numeric=False):
    s = str(value or "").strip()
    if not s:
        return None

    m = re.search(r"\brs\s*[:#_-]?\s*(\d+)\b", s, re.I)
    if m:
        return f"rs{int(m.group(1))}"

    if allow_bare_numeric:
        if re.fullmatch(r"\d+(?:\.0+)?", s):
            return f"rs{int(float(s))}"

    return None


def find_col(headers, candidates):
    lookup = {norm_header(h): i for i, h in enumerate(headers) if h is not None}
    for c in candidates:
        k = norm_header(c)
        if k in lookup:
            return lookup[k]
    return None


def cell(row, idx):
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def canonical_variant_id_from_variation_id(raw):
    """ODDB v1 canonical variant key derived from ClinVar VariationID."""
    s = str(raw or "").strip()
    if not s:
        return None
    if not re.fullmatch(r"\d+(?:\.0+)?", s):
        return None
    n = str(int(float(s)))
    return f"ODDB-VAR-CV{n}"


def main():
    # Parser self-tests.
    assert normalize_rsid("rs80358293") == "rs80358293"
    assert normalize_rsid("RS80358293") == "rs80358293"
    assert normalize_rsid(
        "https://www.ncbi.nlm.nih.gov/snp/rs80358293"
    ) == "rs80358293"
    assert normalize_rsid(
        "80358293", allow_bare_numeric=True
    ) == "rs80358293"

    ap = argparse.ArgumentParser(
        description=(
            "Read-only audit bridging legacy ODDB rsIDs through the frozen "
            "ClinVar staging TSV to canonical ODDB v1 variants and DVA pairs."
        )
    )
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--clinvar-staging", required=True)
    ap.add_argument(
        "--report-dir",
        default="reports/legacy_variant_clinvar_bridge_v0.5.6",
    )
    ap.add_argument("--host", default=os.getenv("MYSQL_HOST", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(os.getenv("MYSQL_PORT", "3306")))
    ap.add_argument("--user", default=os.getenv("MYSQL_USER", "root"))
    ap.add_argument("--password", default=os.getenv("MYSQL_PASSWORD", ""))
    ap.add_argument("--database", default=os.getenv("MYSQL_DATABASE", "oddb_v1"))
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    staging = Path(args.clinvar_staging)
    report_dir = Path(args.report_dir)

    if not xlsx.is_file():
        raise SystemExit(f"Legacy workbook not found: {xlsx}")
    if not staging.is_file():
        raise SystemExit(f"ClinVar staging file not found: {staging}")

    report_dir.mkdir(parents=True, exist_ok=True)

    # ------------------------------------------------------------------
    # 1. Read legacy workbook once and collect rows / target rsIDs.
    # ------------------------------------------------------------------
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = list(next(ws.iter_rows(values_only=True)))

    did_i = find_col(headers, ["DiseaseID"])
    dname_i = find_col(headers, ["Ocular.disease", "Ocular disease"])
    category_i = find_col(headers, ["Disease.category", "Disease category"])
    snpid_i = find_col(headers, ["SNPID", "SNP ID"])
    chr_i = find_col(headers, ["Chromosome"])
    pos_i = find_col(headers, ["Position.hg38", "Position hg38"])
    allele_i = find_col(headers, ["Major.allele", "Major allele"])
    detail_i = find_col(headers, ["variant.Detail", "variant detail"])
    type_i = find_col(headers, ["variant.type", "variant type"])
    protein_i = find_col(headers, ["protein"])
    polyphen_i = find_col(headers, ["PolyPhen"])
    cadd_i = find_col(headers, ["CADD"])
    sift_i = find_col(headers, ["SIFT"])
    gerp_i = find_col(headers, ["GERP"])
    gene_i = find_col(headers, ["Gene"])
    ens_i = find_col(headers, ["Ensembl"])
    dbsnp_i = find_col(headers, ["dbsnp", "dbSNP"])
    gnomad_i = find_col(headers, ["Gnomad", "gnomAD"])
    doi_i = find_col(headers, ["DOIs", "DOI"])

    legacy_rows = []
    target_rsids = set()
    variant_type_counts_all = Counter()

    for row_no, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True), start=2
    ):
        snpid = cell(row, snpid_i)
        dbsnp_url = cell(row, dbsnp_i)
        rsid = (
            normalize_rsid(snpid)
            or normalize_rsid(dbsnp_url)
        )

        vtype = cell(row, type_i)
        if vtype:
            variant_type_counts_all[vtype] += 1

        rec = {
            "Source_Row": row_no,
            "Legacy_DiseaseID": normalize_oddb_id(cell(row, did_i)) or "",
            "Legacy_Disease_Name": cell(row, dname_i),
            "Disease_Category": cell(row, category_i),
            "SNPID": snpid,
            "rsID": rsid or "",
            "Chromosome": cell(row, chr_i),
            "Position_hg38": cell(row, pos_i),
            "Major_Allele": cell(row, allele_i),
            "Variant_Detail": cell(row, detail_i),
            "Variant_Type": vtype,
            "Protein": cell(row, protein_i),
            "PolyPhen": cell(row, polyphen_i),
            "CADD": cell(row, cadd_i),
            "SIFT": cell(row, sift_i),
            "GERP": cell(row, gerp_i),
            "Gene": cell(row, gene_i),
            "Ensembl_URL": cell(row, ens_i),
            "dbSNP_URL": dbsnp_url,
            "gnomAD_URL": cell(row, gnomad_i),
            "DOIs": cell(row, doi_i),
        }
        legacy_rows.append(rec)
        if rsid:
            target_rsids.add(rsid)

    observed_types = set(variant_type_counts_all)
    missing_expected_types = [
        x for x in EXPECTED_VARIANT_TYPES if x not in observed_types
    ]
    unexpected_types = [
        x for x in sorted(observed_types) if x not in EXPECTED_VARIANT_TYPES
    ]

    # ------------------------------------------------------------------
    # 2. Stream frozen ClinVar staging once, collecting only target rsIDs.
    # ------------------------------------------------------------------
    rsid_to_variation_ids = defaultdict(set)
    variation_meta = {}
    staging_rows_scanned = 0
    staging_target_rows = 0

    with staging.open(
        "r", encoding="utf-8", errors="replace", newline=""
    ) as f:
        reader = csv.DictReader(f, delimiter="\t")

        required = {
            "RS# (dbSNP)", "VariationID", "GeneSymbol", "HGNC_ID",
            "Chromosome", "Start", "ReferenceAlleleVCF",
            "AlternateAlleleVCF"
        }
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise SystemExit(
                "ClinVar staging file is missing required columns: "
                + ", ".join(sorted(missing))
            )

        for row in reader:
            staging_rows_scanned += 1
            rsid = normalize_rsid(
                row.get("RS# (dbSNP)"),
                allow_bare_numeric=True
            )
            if not rsid or rsid not in target_rsids:
                continue

            staging_target_rows += 1
            vid = str(row.get("VariationID") or "").strip()
            if not vid:
                continue

            rsid_to_variation_ids[rsid].add(vid)

            if vid not in variation_meta:
                variation_meta[vid] = {
                    "ClinVar_GeneSymbol": row.get("GeneSymbol", ""),
                    "ClinVar_HGNC_ID": row.get("HGNC_ID", ""),
                    "ClinVar_Chromosome": row.get("Chromosome", ""),
                    "ClinVar_Start": row.get("Start", ""),
                    "ClinVar_ReferenceAlleleVCF": row.get(
                        "ReferenceAlleleVCF", ""
                    ),
                    "ClinVar_AlternateAlleleVCF": row.get(
                        "AlternateAlleleVCF", ""
                    ),
                    "ClinVar_ClinicalSignificance": row.get(
                        "ClinicalSignificance", ""
                    ),
                    "ClinVar_ReviewStatus": row.get("ReviewStatus", ""),
                    "ClinVar_RCVaccession": row.get("RCVaccession", ""),
                    "ClinVar_PhenotypeIDS": row.get("PhenotypeIDS", ""),
                    "ClinVar_PhenotypeList": row.get("PhenotypeList", ""),
                }

    # ------------------------------------------------------------------
    # 3. Read current disease aliases and canonical variant IDs from MySQL.
    # ------------------------------------------------------------------
    conn = pymysql.connect(
        host=args.host,
        port=args.port,
        user=args.user,
        password=args.password,
        database=args.database,
        charset="utf8mb4",
        autocommit=False,
        cursorclass=pymysql.cursors.DictCursor,
    )

    try:
        with conn.cursor() as cur:
            cur.execute("SELECT disease_id, disease_name FROM disease")
            diseases = cur.fetchall()

            current_disease_ids = {r["disease_id"] for r in diseases}
            by_name = defaultdict(set)
            for r in diseases:
                by_name[norm_name(r["disease_name"])].add(r["disease_id"])

            cur.execute("SHOW COLUMNS FROM disease_alias")
            alias_cols = [r["Field"] for r in cur.fetchall()]
            alias_lookup = {norm_header(c): c for c in alias_cols}
            alias_col = alias_lookup.get("alias")
            source_id_col = alias_lookup.get("sourceid")

            old_to_current = {}

            if alias_col:
                cur.execute(
                    f"SELECT disease_id, `{alias_col}` AS alias_text "
                    f"FROM disease_alias "
                    f"WHERE `{alias_col}` IS NOT NULL "
                    f"AND TRIM(`{alias_col}`)<>''"
                )
                for r in cur.fetchall():
                    by_name[norm_name(r["alias_text"])].add(r["disease_id"])

            if source_id_col:
                cur.execute(
                    f"SELECT disease_id, `{source_id_col}` AS source_id "
                    f"FROM disease_alias "
                    f"WHERE `{source_id_col}` IS NOT NULL "
                    f"AND TRIM(`{source_id_col}`)<>''"
                )
                for r in cur.fetchall():
                    sid = normalize_oddb_id(r["source_id"])
                    if sid:
                        old_to_current[sid] = r["disease_id"]

            # Known retired exact alias.
            old_to_current.setdefault("ODDB-76", "ODDB-225")

            cur.execute("SHOW COLUMNS FROM variant")
            variant_cols = [r["Field"] for r in cur.fetchall()]
            variant_lookup = {norm_header(c): c for c in variant_cols}
            variant_id_col = variant_lookup.get("variantid")
            if not variant_id_col:
                raise SystemExit("Could not detect variant.variant_id")

            # Canonical ODDB v1 variant IDs are deterministic:
            # ClinVar VariationID N -> ODDB-VAR-CV<N>.
            expected_canonical_ids = {}
            for vid in variation_meta:
                cid = canonical_variant_id_from_variation_id(vid)
                if cid:
                    expected_canonical_ids[vid] = cid

            canonical_variant_ids = set()

            expected_values = sorted(set(expected_canonical_ids.values()))
            for i in range(0, len(expected_values), 1000):
                chunk = expected_values[i:i + 1000]
                placeholders = ",".join(["%s"] * len(chunk))
                cur.execute(
                    f"SELECT `{variant_id_col}` AS variant_id "
                    f"FROM variant "
                    f"WHERE `{variant_id_col}` IN ({placeholders})",
                    chunk,
                )
                for r in cur.fetchall():
                    canonical_variant_ids.add(str(r["variant_id"]))

            variation_to_canonical = {
                vid: cid
                for vid, cid in expected_canonical_ids.items()
                if cid in canonical_variant_ids
            }
            ambiguous_canonical_forms = {}

            # Print samples for diagnosis.
            cur.execute(
                f"SELECT `{variant_id_col}` AS variant_id "
                f"FROM variant ORDER BY `{variant_id_col}` LIMIT 10"
            )
            canonical_variant_samples = [
                str(r["variant_id"]) for r in cur.fetchall()
            ]

            cur.execute("""
                SELECT disease_id, variant_id
                FROM disease_variant_association
            """)
            current_dva_pairs = {
                (str(r["disease_id"]), str(r["variant_id"]))
                for r in cur.fetchall()
            }

        # --------------------------------------------------------------
        # 4. Resolve every legacy row.
        # --------------------------------------------------------------
        disease_resolution = Counter()
        rs_bridge_status = Counter()
        canonical_variant_status = Counter()
        pair_status = Counter()
        safe_type_counts = Counter()

        safe_rows = []
        unresolved_disease_rows = []
        unmatched_rsid_rows = []
        ambiguous_rsid_rows = []
        canonical_unmatched_rows = []
        pair_not_in_dva_rows = []

        safe_variant_categories = defaultdict(set)
        safe_variants = set()
        safe_diseases = set()

        for rec in legacy_rows:
            # Disease resolution.
            disease_hits = set()
            disease_methods = []

            legacy_did = rec["Legacy_DiseaseID"]
            dname = rec["Legacy_Disease_Name"]

            if legacy_did in current_disease_ids:
                disease_hits.add(legacy_did)
                disease_methods.append("current DiseaseID")

            if legacy_did in old_to_current:
                disease_hits.add(old_to_current[legacy_did])
                disease_methods.append("retired/source DiseaseID")

            if dname:
                hits = by_name.get(norm_name(dname), set())
                if hits:
                    disease_hits.update(hits)
                    disease_methods.append("exact disease name/alias")

            if len(disease_hits) != 1:
                disease_resolution[
                    "ambiguous" if len(disease_hits) > 1 else "unresolved"
                ] += 1
                unresolved_disease_rows.append({
                    **rec,
                    "Candidate_Current_DiseaseIDs": "|".join(
                        sorted(disease_hits)
                    )
                })
                continue

            current_did = next(iter(disease_hits))
            disease_resolution["resolved"] += 1

            rsid = rec["rsID"]
            if not rsid:
                rs_bridge_status["no_rsid"] += 1
                unmatched_rsid_rows.append({
                    **rec,
                    "Current_DiseaseID": current_did,
                    "Reason": "no rsID extracted from legacy row",
                })
                continue

            vids = rsid_to_variation_ids.get(rsid, set())
            if not vids:
                rs_bridge_status["not_in_staging"] += 1
                unmatched_rsid_rows.append({
                    **rec,
                    "Current_DiseaseID": current_did,
                    "Reason": "rsID not found in frozen ClinVar staging",
                })
                continue

            if len(vids) > 1:
                # Conservative disambiguation using only already-supported
                # current disease_variant_association pairs.
                dva_supported = []
                for candidate_vid in sorted(vids):
                    candidate_canonical = variation_to_canonical.get(candidate_vid)
                    if candidate_canonical and (
                        str(current_did), str(candidate_canonical)
                    ) in current_dva_pairs:
                        dva_supported.append((candidate_vid, candidate_canonical))

                if len(dva_supported) == 1:
                    rs_bridge_status["resolved_ambiguous_by_current_DVA"] += 1
                    clinvar_vid, canonical_vid = dva_supported[0]
                    pair_present = True
                else:
                    rs_bridge_status["ambiguous_multiple_variationids"] += 1
                    ambiguous_rsid_rows.append({
                        **rec,
                        "Current_DiseaseID": current_did,
                        "ClinVar_VariationIDs": "|".join(sorted(vids)),
                        "Current_DVA_Matching_Candidates": "|".join(
                            f"{v}:{c}" for v, c in dva_supported
                        ),
                    })
                    continue
            else:
                rs_bridge_status["unique_variationid"] += 1
                clinvar_vid = next(iter(vids))
                canonical_vid = variation_to_canonical.get(clinvar_vid)
                pair_present = None

            if not canonical_vid:
                canonical_variant_status["not_in_canonical_variant"] += 1
                canonical_unmatched_rows.append({
                    **rec,
                    "Current_DiseaseID": current_did,
                    "ClinVar_VariationID": clinvar_vid,
                    "Reason": (
                        "VariationID did not resolve to canonical variant_id"
                    ),
                })
                continue

            canonical_variant_status["resolved"] += 1

            if pair_present is None:
                pair_present = (
                    str(current_did), str(canonical_vid)
                ) in current_dva_pairs

            pair_status[
                "current_DVA_match" if pair_present
                else "not_in_current_DVA"
            ] += 1

            outrec = {
                **rec,
                "Current_DiseaseID": current_did,
                "Disease_Mapping_Method": "+".join(
                    sorted(set(disease_methods))
                ),
                "ClinVar_VariationID": clinvar_vid,
                "Current_VariantID": canonical_vid,
                "Current_DVA_Pair": "yes" if pair_present else "no",
                **variation_meta.get(clinvar_vid, {}),
            }

            if pair_present:
                safe_rows.append(outrec)
                safe_variants.add(canonical_vid)
                safe_diseases.add(current_did)
                if rec["Variant_Type"]:
                    safe_variant_categories[canonical_vid].add(
                        rec["Variant_Type"]
                    )
                    safe_type_counts[rec["Variant_Type"]] += 1
            else:
                pair_not_in_dva_rows.append(outrec)

        matched_staging_rsids = set(rsid_to_variation_ids)
        unmatched_unique_rsids = target_rsids - matched_staging_rsids
        ambiguous_unique_rsids = {
            rs for rs, vids in rsid_to_variation_ids.items()
            if len(vids) > 1
        }
        one_to_one_unique_rsids = {
            rs for rs, vids in rsid_to_variation_ids.items()
            if len(vids) == 1
        }

        safe_multi_category_variants = {
            vid: cats
            for vid, cats in safe_variant_categories.items()
            if len(cats) > 1
        }

        # --------------------------------------------------------------
        # 5. Print audit summary.
        # --------------------------------------------------------------
        print("ODDB LEGACY VARIANT → CLINVAR → CANONICAL AUDIT v0.5.6")
        print("=" * 96)
        print(f"Legacy workbook rows: {len(legacy_rows):,}")
        print(f"Legacy unique rsIDs: {len(target_rsids):,}")
        print(f"Legacy variant.type categories: {len(observed_types):,}")

        print("\nClinVar staging bridge:")
        print(f"Staging rows scanned: {staging_rows_scanned:,}")
        print(f"Staging rows matching target rsIDs: {staging_target_rows:,}")
        print(f"Unique legacy rsIDs found in staging: "
              f"{len(matched_staging_rsids):,}")
        print(f"  one rsID -> one VariationID: "
              f"{len(one_to_one_unique_rsids):,}")
        print(f"  one rsID -> multiple VariationIDs: "
              f"{len(ambiguous_unique_rsids):,}")
        print(f"Legacy rsIDs absent from staging: "
              f"{len(unmatched_unique_rsids):,}")

        print("\nCanonical variant bridge:")
        print(f"ClinVar VariationIDs collected: {len(variation_meta):,}")
        print(f"VariationIDs resolved to canonical variant_id: "
              f"{len(variation_to_canonical):,}")
        print(f"VariationIDs with ambiguous canonical forms: "
              f"{len(ambiguous_canonical_forms):,}")
        print("Sample canonical variant_ids: "
              + ", ".join(canonical_variant_samples[:10]))

        print("\nDisease resolution:")
        for k in ["resolved", "unresolved", "ambiguous"]:
            print(f"{k:24s} {disease_resolution[k]:7d}")

        print("\nLegacy rsID → ClinVar VariationID row status:")
        for k in [
            "unique_variationid",
            "resolved_ambiguous_by_current_DVA",
            "not_in_staging",
            "ambiguous_multiple_variationids",
            "no_rsid",
        ]:
            print(f"{k:36s} {rs_bridge_status[k]:7d}")

        print("\nClinVar VariationID → canonical variant row status:")
        for k in ["resolved", "not_in_canonical_variant"]:
            print(f"{k:36s} {canonical_variant_status[k]:7d}")

        print("\nDisease–variant pair verification:")
        for k in ["current_DVA_match", "not_in_current_DVA"]:
            print(f"{k:36s} {pair_status[k]:7d}")

        print("\nSAFE consequence layer candidates:")
        print(f"Safe legacy rows with current DVA match: {len(safe_rows):,}")
        print(f"Unique current diseases represented: {len(safe_diseases):,}")
        print(f"Unique current variants represented: {len(safe_variants):,}")
        print(f"Unique variant/category memberships: "
              f"{sum(len(x) for x in safe_variant_categories.values()):,}")
        print(f"Variants with >1 legacy variant.type category: "
              f"{len(safe_multi_category_variants):,}")

        print("\nSafe row counts by variant.type:")
        for value, n in safe_type_counts.most_common():
            print(f"{n:7d}  {value}")

        print("\nExpected taxonomy check:")
        print(f"Missing expected types: {missing_expected_types or 'none'}")
        print(f"Unexpected types: {unexpected_types or 'none'}")

        # --------------------------------------------------------------
        # 6. Reports.
        # --------------------------------------------------------------
        base_fields = [
            "Source_Row", "Legacy_DiseaseID", "Legacy_Disease_Name",
            "Current_DiseaseID", "Disease_Mapping_Method",
            "Disease_Category", "SNPID", "rsID",
            "ClinVar_VariationID", "Current_VariantID",
            "Current_DVA_Pair",
            "Chromosome", "Position_hg38", "Major_Allele",
            "Variant_Detail", "Variant_Type",
            "Protein", "PolyPhen", "CADD", "SIFT", "GERP", "Gene",
            "Ensembl_URL", "dbSNP_URL", "gnomAD_URL", "DOIs",
            "ClinVar_GeneSymbol", "ClinVar_HGNC_ID",
            "ClinVar_Chromosome", "ClinVar_Start",
            "ClinVar_ReferenceAlleleVCF", "ClinVar_AlternateAlleleVCF",
            "ClinVar_ClinicalSignificance", "ClinVar_ReviewStatus",
            "ClinVar_RCVaccession", "ClinVar_PhenotypeIDS",
            "ClinVar_PhenotypeList",
        ]

        def write_dicts(path, rows, fields):
            with path.open("w", encoding="utf-8", newline="") as f:
                w = csv.DictWriter(
                    f, fieldnames=fields, delimiter="\t",
                    extrasaction="ignore"
                )
                w.writeheader()
                w.writerows(rows)

        write_dicts(
            report_dir / "safe_mapped_rows_current_dva.tsv",
            safe_rows,
            base_fields,
        )

        write_dicts(
            report_dir / "mapped_rows_not_in_current_dva.tsv",
            pair_not_in_dva_rows,
            base_fields,
        )

        unresolved_fields = [
            "Source_Row", "Legacy_DiseaseID", "Legacy_Disease_Name",
            "SNPID", "rsID", "Variant_Type",
            "Candidate_Current_DiseaseIDs",
        ]
        write_dicts(
            report_dir / "unresolved_disease_rows.tsv",
            unresolved_disease_rows,
            unresolved_fields,
        )

        unmatched_fields = [
            "Source_Row", "Legacy_DiseaseID", "Legacy_Disease_Name",
            "Current_DiseaseID", "SNPID", "rsID",
            "Variant_Type", "Reason",
        ]
        write_dicts(
            report_dir / "unmatched_rsid_rows.tsv",
            unmatched_rsid_rows,
            unmatched_fields,
        )

        ambiguous_fields = [
            "Source_Row", "Legacy_DiseaseID", "Legacy_Disease_Name",
            "Current_DiseaseID", "SNPID", "rsID",
            "Variant_Type", "ClinVar_VariationIDs",
            "Current_DVA_Matching_Candidates",
        ]
        write_dicts(
            report_dir / "ambiguous_rsid_rows.tsv",
            ambiguous_rsid_rows,
            ambiguous_fields,
        )

        canonical_unmatched_fields = [
            "Source_Row", "Legacy_DiseaseID", "Legacy_Disease_Name",
            "Current_DiseaseID", "SNPID", "rsID",
            "Variant_Type", "ClinVar_VariationID", "Reason",
        ]
        write_dicts(
            report_dir / "clinvar_variationid_not_in_canonical_variant.tsv",
            canonical_unmatched_rows,
            canonical_unmatched_fields,
        )

        with (
            report_dir / "safe_variant_consequence_memberships.tsv"
        ).open("w", encoding="utf-8", newline="") as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow([
                "Current_VariantID", "Variant_Type",
                "Legacy_Supporting_Row_Count"
            ])

            row_support = Counter()
            for r in safe_rows:
                row_support[
                    (r["Current_VariantID"], r["Variant_Type"])
                ] += 1

            for (vid, vtype), n in sorted(row_support.items()):
                w.writerow([vid, vtype, n])

        with (
            report_dir / "rsid_to_clinvar_variationid.tsv"
        ).open("w", encoding="utf-8", newline="") as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow(["rsID", "ClinVar_VariationID", "Mapping_Status"])
            for rsid in sorted(target_rsids):
                vids = sorted(rsid_to_variation_ids.get(rsid, set()))
                if not vids:
                    w.writerow([rsid, "", "not_in_staging"])
                elif len(vids) == 1:
                    w.writerow([rsid, vids[0], "unique"])
                else:
                    for vid in vids:
                        w.writerow([rsid, vid, "ambiguous_multiple_variationids"])

        with (
            report_dir / "safe_variant_type_counts.tsv"
        ).open("w", encoding="utf-8", newline="") as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow(["Variant_Type", "Safe_Legacy_Row_Count"])
            for value, n in safe_type_counts.most_common():
                w.writerow([value, n])

        summary = {
            "version": "v0.5.6",
            "legacy_rows": len(legacy_rows),
            "legacy_unique_rsids": len(target_rsids),
            "legacy_variant_type_categories": len(observed_types),
            "staging_rows_scanned": staging_rows_scanned,
            "staging_target_rows": staging_target_rows,
            "legacy_rsids_found_in_staging": len(matched_staging_rsids),
            "one_to_one_rsids": len(one_to_one_unique_rsids),
            "ambiguous_rsids": len(ambiguous_unique_rsids),
            "unmatched_rsids": len(unmatched_unique_rsids),
            "clinvar_variationids_collected": len(variation_meta),
            "variationids_resolved_to_canonical": len(variation_to_canonical),
            "rows_resolved_from_ambiguous_rsid_by_current_DVA":
                rs_bridge_status["resolved_ambiguous_by_current_DVA"],
            "safe_current_dva_rows": len(safe_rows),
            "safe_unique_current_diseases": len(safe_diseases),
            "safe_unique_current_variants": len(safe_variants),
            "safe_variant_category_memberships": sum(
                len(x) for x in safe_variant_categories.values()
            ),
            "safe_multi_category_variants": len(
                safe_multi_category_variants
            ),
            "missing_expected_variant_types": missing_expected_types,
            "unexpected_variant_types": unexpected_types,
        }

        with (
            report_dir / "audit_summary.json"
        ).open("w", encoding="utf-8") as f:
            json.dump(summary, f, indent=2, ensure_ascii=False)

        print("\nReports:")
        for name in [
            "safe_mapped_rows_current_dva.tsv",
            "mapped_rows_not_in_current_dva.tsv",
            "unresolved_disease_rows.tsv",
            "unmatched_rsid_rows.tsv",
            "ambiguous_rsid_rows.tsv",
            "clinvar_variationid_not_in_canonical_variant.tsv",
            "safe_variant_consequence_memberships.tsv",
            "rsid_to_clinvar_variationid.tsv",
            "safe_variant_type_counts.tsv",
            "audit_summary.json",
        ]:
            print(report_dir / name)

        print("\nREAD-ONLY: MySQL, frozen releases, staging data, "
              "and legacy workbook were not modified.")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
