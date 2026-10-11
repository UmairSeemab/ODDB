#!/usr/bin/env python3
from __future__ import annotations

import argparse
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pymysql


CATEGORY_MAP = {
    "autosomal dominant": "INH-01",
    "autosomal recessive": "INH-02",
    "isolated cases": "INH-03",
    "mitochondrial": "INH-04",
    "other": "INH-05",
    "x-linked dominant": "INH-06",
    "x-linked recessive": "INH-07",
}

CATEGORY_NAMES = {
    "INH-01": "Autosomal dominant",
    "INH-02": "Autosomal recessive",
    "INH-03": "Isolated cases",
    "INH-04": "Mitochondrial",
    "INH-05": "Other",
    "INH-06": "X-linked dominant",
    "INH-07": "X-linked recessive",
}


def norm_header(x):
    x = str(x or "").strip().lower()
    x = re.sub(r"[^a-z0-9]+", "", x)
    return x


def norm_name(x):
    x = unicodedata.normalize("NFKD", str(x or ""))
    x = "".join(ch for ch in x if not unicodedata.combining(ch))
    x = x.upper().strip()
    x = x.replace("’", "'").replace("–", "-").replace("—", "-")
    x = re.sub(r"\s+", " ", x)
    x = re.sub(r"\s*[-_/]\s*", " ", x)
    x = re.sub(r"[^A-Z0-9]+", " ", x)
    return re.sub(r"\s+", " ", x).strip()


def norm_mode(x):
    x = str(x or "").strip().lower()
    x = x.replace("_", " ").replace("–", "-").replace("—", "-")
    x = re.sub(r"\s+", " ", x)
    x = re.sub(r"^\d+\s*[.)-]?\s*", "", x)
    synonyms = {
        "autosomal dominant inheritance": "autosomal dominant",
        "autosomal recessive inheritance": "autosomal recessive",
        "isolated case": "isolated cases",
        "isolated": "isolated cases",
        "sporadic": "isolated cases",
        "maternal": "mitochondrial",
        "maternal inheritance": "mitochondrial",
        "x linked dominant": "x-linked dominant",
        "x linked recessive": "x-linked recessive",
        "x-linked": "other",
        "x linked": "other",
        "unknown": "other",
        "others": "other",
    }
    return synonyms.get(x, x)


def find_column(headers, candidates):
    lookup = {norm_header(h): i for i, h in enumerate(headers) if h is not None}
    for c in candidates:
        k = norm_header(c)
        if k in lookup:
            return lookup[k]
    return None


def detect_sheet(wb):
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration:
            continue
        inh = find_column(headers, [
            "MODE.OF.INHERITANCE", "MODE OF INHERITANCE", "Inheritance",
            "Inheritance mode", "Mode_of_Inheritance"
        ])
        if inh is not None:
            return ws, headers
    raise SystemExit("No worksheet with a MODE.OF.INHERITANCE/Inheritance column was found.")


def main():
    ap = argparse.ArgumentParser(
        description="Import the original ODDB inheritance taxonomy from Inheritance_Based_file.xlsx."
    )
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--apply", action="store_true",
                    help="Write mappings to MySQL. Without this flag the script is preview-only.")
    ap.add_argument("--report-dir", default="reports/web_taxonomy_inheritance_legacy_v0.3.6")
    ap.add_argument("--host", default=os.getenv("MYSQL_HOST", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(os.getenv("MYSQL_PORT", "3306")))
    ap.add_argument("--user", default=os.getenv("MYSQL_USER", "root"))
    ap.add_argument("--password", default=os.getenv("MYSQL_PASSWORD", ""))
    ap.add_argument("--database", default=os.getenv("MYSQL_DATABASE", "oddb_v1"))
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    if not xlsx.is_file():
        raise SystemExit(f"Workbook not found: {xlsx}")

    report_dir = Path(args.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws, headers = detect_sheet(wb)

    inh_i = find_column(headers, [
        "MODE.OF.INHERITANCE", "MODE OF INHERITANCE", "Inheritance",
        "Inheritance mode", "Mode_of_Inheritance"
    ])
    did_i = find_column(headers, [
        "DiseaseID", "Disease ID", "Disease_ID", "ODDB_DiseaseID", "ODDB Disease ID"
    ])
    dname_i = find_column(headers, [
        "Ocular.disease", "Ocular disease", "Disease", "Disease name",
        "Disease_Name", "OcularDisease"
    ])

    print("ODDB v1 LEGACY INHERITANCE XLSX IMPORTER v0.3.6")
    print("=" * 76)
    print(f"Workbook: {xlsx}")
    print(f"Worksheet: {ws.title}")
    print(f"Headers: {headers}")
    print(f"Inheritance column: {headers[inh_i] if inh_i is not None else None}")
    print(f"Disease ID column: {headers[did_i] if did_i is not None else None}")
    print(f"Disease name column: {headers[dname_i] if dname_i is not None else None}")

    if did_i is None and dname_i is None:
        raise SystemExit("No disease ID or disease-name column could be detected.")

    conn = pymysql.connect(
        host=args.host, port=args.port, user=args.user, password=args.password,
        database=args.database, charset="utf8mb4", autocommit=False,
        cursorclass=pymysql.cursors.DictCursor,
    )

    try:
        with conn.cursor() as cur:
            cur.execute("SELECT disease_id,disease_name FROM disease")
            diseases = cur.fetchall()
            by_id = {r["disease_id"]: r["disease_id"] for r in diseases}
            by_name = defaultdict(list)
            for r in diseases:
                by_name[norm_name(r["disease_name"])].append(r["disease_id"])

            # Include aliases as exact-name lookup support.
            cur.execute("""
                SELECT disease_id,alias_name
                FROM disease_alias
                WHERE alias_name IS NOT NULL AND alias_name <> ''
            """)
            for r in cur.fetchall():
                by_name[norm_name(r["alias_name"])].append(r["disease_id"])

            raw_mode_counts = Counter()
            mapped = defaultdict(lambda: defaultdict(set))
            unmapped_modes = Counter()
            unresolved_rows = []
            ambiguous_rows = []
            duplicate_source_rows = 0
            seen_source = set()

            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                mode_raw = row[inh_i] if inh_i is not None and inh_i < len(row) else None
                if mode_raw is None or str(mode_raw).strip() == "":
                    continue

                mode = norm_mode(mode_raw)
                raw_mode_counts[str(mode_raw).strip()] += 1
                cid = CATEGORY_MAP.get(mode)
                if not cid:
                    unmapped_modes[str(mode_raw).strip()] += 1
                    continue

                did_raw = str(row[did_i]).strip() if did_i is not None and did_i < len(row) and row[did_i] is not None else ""
                dname_raw = str(row[dname_i]).strip() if dname_i is not None and dname_i < len(row) and row[dname_i] is not None else ""

                source_key = (did_raw, dname_raw, mode)
                if source_key in seen_source:
                    duplicate_source_rows += 1
                seen_source.add(source_key)

                resolved = None
                basis = None

                if did_raw and did_raw in by_id:
                    resolved = did_raw
                    basis = f"legacy DiseaseID={did_raw}"
                elif dname_raw:
                    hits = sorted(set(by_name.get(norm_name(dname_raw), [])))
                    if len(hits) == 1:
                        resolved = hits[0]
                        basis = f"legacy exact disease name={dname_raw}"
                    elif len(hits) > 1:
                        ambiguous_rows.append((row_num, did_raw, dname_raw, str(mode_raw), "|".join(hits)))
                        continue

                if not resolved:
                    unresolved_rows.append((row_num, did_raw, dname_raw, str(mode_raw)))
                    continue

                mapped[resolved][cid].add(f"{basis}; MODE.OF.INHERITANCE={str(mode_raw).strip()}")

            membership_count = sum(len(v) for v in mapped.values())

            print(f"\nSource rows with inheritance: {sum(raw_mode_counts.values()):,}")
            print(f"Unique mapped current diseases: {len(mapped):,}")
            print(f"Mapped disease/category memberships: {membership_count:,}")
            print(f"Duplicate source rows: {duplicate_source_rows:,}")
            print(f"Unresolved source rows: {len(unresolved_rows):,}")
            print(f"Ambiguous source rows: {len(ambiguous_rows):,}")

            print("\nRaw inheritance values:")
            for value, n in raw_mode_counts.most_common():
                print(f"{n:6d}  {value}")

            if unmapped_modes:
                print("\nUNMAPPED inheritance labels:")
                for value, n in unmapped_modes.most_common():
                    print(f"{n:6d}  {value}")

            counts = Counter()
            for did, cats in mapped.items():
                for cid in cats:
                    counts[cid] += 1

            print("\nPreview category counts:")
            for cid in CATEGORY_NAMES:
                print(f"{CATEGORY_NAMES[cid]:28s} {counts[cid]:4d}")

            # Reports
            with (report_dir / "unresolved_rows.tsv").open("w", encoding="utf-8") as f:
                f.write("Row\tDiseaseID\tDisease_Name\tInheritance\n")
                for r in unresolved_rows:
                    f.write("\t".join(map(str, r)) + "\n")

            with (report_dir / "ambiguous_rows.tsv").open("w", encoding="utf-8") as f:
                f.write("Row\tDiseaseID\tDisease_Name\tInheritance\tCandidate_DiseaseIDs\n")
                for r in ambiguous_rows:
                    f.write("\t".join(map(str, r)) + "\n")

            with (report_dir / "mapped_memberships.tsv").open("w", encoding="utf-8") as f:
                f.write("DiseaseID\tCategory_ID\tCategory_Name\tMembership_Basis\n")
                for did in sorted(mapped):
                    for cid in sorted(mapped[did]):
                        f.write(
                            f"{did}\t{cid}\t{CATEGORY_NAMES[cid]}\t"
                            + " || ".join(sorted(mapped[did][cid]))
                            + "\n"
                        )

            if not args.apply:
                print("\nPREVIEW ONLY — database was not changed.")
                print("Review the counts/reports, then rerun with --apply if they look correct.")
                print(f"Reports: {report_dir}")
                return

            # Replace ONLY inheritance web taxonomy memberships.
            cur.execute("""
                DELETE m
                FROM disease_category_membership m
                JOIN disease_category c ON c.category_id=m.category_id
                WHERE c.category_system='inheritance'
            """)

            inserted = 0
            for did in sorted(mapped):
                for cid in sorted(mapped[did]):
                    basis = " || ".join(sorted(mapped[did][cid]))
                    cur.execute("""
                        INSERT INTO disease_category_membership
                          (disease_id,category_id,membership_basis,source,curation_status)
                        VALUES (%s,%s,%s,'ODDB legacy Inheritance_Based_file.xlsx','legacy_curated')
                    """, (did, cid, basis[:255]))
                    inserted += 1

            conn.commit()

            print(f"\nAPPLIED: {inserted:,} inheritance memberships written to MySQL.")

            cur.execute("""
                SELECT c.category_name, COUNT(DISTINCT m.disease_id) AS n
                FROM disease_category c
                LEFT JOIN disease_category_membership m ON m.category_id=c.category_id
                WHERE c.category_system='inheritance'
                GROUP BY c.category_id,c.category_name,c.display_order
                ORDER BY c.display_order
            """)
            print("\nDatabase category counts after import:")
            for r in cur.fetchall():
                print(f"{r['category_name']:28s} {r['n']:4d}")

            cur.execute("""
                SELECT COUNT(DISTINCT m.disease_id) AS n
                FROM disease_category_membership m
                JOIN disease_category c ON c.category_id=m.category_id
                WHERE c.category_system='inheritance'
            """)
            covered = cur.fetchone()["n"]
            print(f"\nCurrent diseases represented in inheritance taxonomy: {covered} / {len(diseases)}")
            print(f"Reports: {report_dir}")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
