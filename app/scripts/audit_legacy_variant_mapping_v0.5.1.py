#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pymysql


EXPECTED_VARIANT_TYPES = [
    "missense variant",
    "inframe deletion",
    "frameshift variant",
    "intron variant",
    "regulatory region variant",
    "intergenic variant",
    "splice region variant",
    "splice donor variant",
    "non coding transcript exon variant",
    "3 prime UTR variant",
    "5 prime UTR variant",
    "stop gained",
    "synonymous variant",
    "TF binding site variant",
    "splice acceptor variant",
    "downstream gene variant",
    "stop lost",
    "upstream gene variant",
    "inframe insertion",
    "protein altering variant",
    "Multiple",
]


def norm_header(x):
    return re.sub(r"[^a-z0-9]+", "", str(x or "").strip().lower())


def norm_name(x):
    x = unicodedata.normalize("NFKD", str(x or ""))
    x = "".join(ch for ch in x if not unicodedata.combining(ch))
    x = x.upper().strip()
    x = x.replace("’", "'").replace("–", "-").replace("—", "-")
    x = re.sub(r"\s*[-_/]\s*", " ", x)
    x = re.sub(r"[^A-Z0-9]+", " ", x)
    return re.sub(r"\s+", " ", x).strip()


def normalize_oddb_id(x):
    s = str(x or "").strip().upper()
    if not s:
        return None
    m = re.fullmatch(r"ODDB[\s_-]*(\d+)(?:\.0+)?", s)
    if m:
        return f"ODDB-{int(m.group(1))}"
    return s


def extract_rsid(*values):
    for value in values:
        s = str(value or "")
        m = re.search(r"\brs\d+\b", s, re.I)
        if m:
            return m.group(0).lower()
    return None


def find_col(headers, candidates):
    lookup = {norm_header(h): i for i, h in enumerate(headers) if h is not None}
    for c in candidates:
        k = norm_header(c)
        if k in lookup:
            return lookup[k]
    return None


def cell(row, idx):
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def main():
    ap = argparse.ArgumentParser(
        description="Audit mapping of legacy variant workbook to current ODDB v1."
    )
    ap.add_argument("--xlsx", required=True)
    ap.add_argument(
        "--report-dir",
        default="reports/legacy_variant_mapping_audit_v0.5.1",
    )
    ap.add_argument("--host", default=os.getenv("MYSQL_HOST", "127.0.0.1"))
    ap.add_argument("--port", type=int, default=int(os.getenv("MYSQL_PORT", "3306")))
    ap.add_argument("--user", default=os.getenv("MYSQL_USER", "root"))
    ap.add_argument("--password", default=os.getenv("MYSQL_PASSWORD", ""))
    ap.add_argument("--database", default=os.getenv("MYSQL_DATABASE", "oddb_v1"))
    args = ap.parse_args()

    xlsx = Path(args.xlsx)
    if not xlsx.is_file():
        raise SystemExit(f"Workbook not found: {xlsx}")

    report_dir = Path(args.report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = list(next(ws.iter_rows(values_only=True)))

    did_i = find_col(headers, ["DiseaseID"])
    dname_i = find_col(headers, ["Ocular.disease", "Ocular disease"])
    category_i = find_col(headers, ["Disease.category", "Disease category"])
    snpid_i = find_col(headers, ["SNPID", "SNP ID"])
    chr_i = find_col(headers, ["Chromosome"])
    pos_i = find_col(headers, ["Position.hg38", "Position hg38"])
    allele_i = find_col(headers, ["Major.allele", "Major allele"])
    detail_i = find_col(headers, ["variant.Detail", "variant detail"])
    type_i = find_col(headers, ["variant.type", "variant type"])
    gene_i = find_col(headers, ["Gene"])
    ens_i = find_col(headers, ["Ensembl"])
    dbsnp_i = find_col(headers, ["dbsnp", "dbSNP"])
    gnomad_i = find_col(headers, ["Gnomad", "gnomAD"])
    doi_i = find_col(headers, ["DOIs", "DOI"])

    conn = pymysql.connect(
        host=args.host,
        port=args.port,
        user=args.user,
        password=args.password,
        database=args.database,
        charset="utf8mb4",
        autocommit=False,
        cursorclass=pymysql.cursors.DictCursor,
    )

    try:
        with conn.cursor() as cur:
            # ---------------------------
            # Current disease resolution
            # ---------------------------
            cur.execute("SELECT disease_id, disease_name FROM disease")
            diseases = cur.fetchall()

            current_ids = {r["disease_id"] for r in diseases}
            by_name = defaultdict(set)
            for r in diseases:
                by_name[norm_name(r["disease_name"])].add(r["disease_id"])

            cur.execute("SHOW COLUMNS FROM disease_alias")
            alias_cols = [r["Field"] for r in cur.fetchall()]
            alias_lookup = {norm_header(c): c for c in alias_cols}

            alias_col = alias_lookup.get("alias")
            source_id_col = alias_lookup.get("sourceid")
            old_to_current = {}

            if alias_col:
                cur.execute(
                    f"SELECT disease_id, `{alias_col}` AS alias_text "
                    f"FROM disease_alias "
                    f"WHERE `{alias_col}` IS NOT NULL AND `{alias_col}`<>''"
                )
                for r in cur.fetchall():
                    by_name[norm_name(r["alias_text"])].add(r["disease_id"])

            if source_id_col:
                cur.execute(
                    f"SELECT disease_id, `{source_id_col}` AS source_id "
                    f"FROM disease_alias "
                    f"WHERE `{source_id_col}` IS NOT NULL AND `{source_id_col}`<>''"
                )
                for r in cur.fetchall():
                    sid = normalize_oddb_id(r["source_id"])
                    if sid:
                        old_to_current[sid] = r["disease_id"]

            old_to_current.setdefault("ODDB-76", "ODDB-225")

            # ---------------------------
            # Current variant rsID index
            # ---------------------------
            cur.execute("SHOW COLUMNS FROM variant")
            variant_cols = [r["Field"] for r in cur.fetchall()]
            variant_lookup = {norm_header(c): c for c in variant_cols}

            rsid_col = None
            for candidate in [
                "rsid", "rs_id", "dbsnp_id", "dbsnp", "rs"
            ]:
                k = norm_header(candidate)
                if k in variant_lookup:
                    rsid_col = variant_lookup[k]
                    break

            variant_id_col = None
            for candidate in [
                "variant_id", "variation_id", "clinvar_variation_id"
            ]:
                k = norm_header(candidate)
                if k in variant_lookup:
                    variant_id_col = variant_lookup[k]
                    break

            if variant_id_col is None:
                raise SystemExit(
                    "Could not detect the canonical variant primary identifier column."
                )

            source_json_col = None
            if "sourcerecordjson" in variant_lookup:
                source_json_col = variant_lookup["sourcerecordjson"]

            rsid_to_variant = defaultdict(set)

            if rsid_col:
                cur.execute(
                    f"SELECT `{variant_id_col}` AS variant_id, "
                    f"`{rsid_col}` AS rsid "
                    f"FROM variant "
                    f"WHERE `{rsid_col}` IS NOT NULL AND `{rsid_col}`<>''"
                )
                for r in cur.fetchall():
                    rs = extract_rsid(r["rsid"])
                    if rs:
                        rsid_to_variant[rs].add(str(r["variant_id"]))
            elif source_json_col:
                # Fallback: parse frozen source JSON. This scans the variant table
                # but keeps mapping read-only and exact.
                cur.execute(
                    f"SELECT `{variant_id_col}` AS variant_id, "
                    f"`{source_json_col}` AS source_record_json "
                    f"FROM variant"
                )
                for r in cur.fetchall():
                    raw = r["source_record_json"]
                    if not raw:
                        continue
                    rs = None
                    try:
                        obj = json.loads(raw)
                        if isinstance(obj, dict):
                            for key in [
                                "RS# (dbSNP)", "RS#",
                                "rsid", "RSID", "dbSNP", "dbsnp"
                            ]:
                                if key in obj and obj[key]:
                                    rs = extract_rsid(obj[key])
                                    if rs:
                                        break
                    except Exception:
                        rs = extract_rsid(raw)

                    if rs:
                        rsid_to_variant[rs].add(str(r["variant_id"]))
            else:
                raise SystemExit(
                    "Variant table has neither a detected rsID column nor source_record_json."
                )

            # Current DVA pairs for semantic verification.
            cur.execute("""
                SELECT disease_id, variant_id
                FROM disease_variant_association
            """)
            current_dva_pairs = {
                (str(r["disease_id"]), str(r["variant_id"]))
                for r in cur.fetchall()
            }

        # ---------------------------
        # Legacy workbook audit
        # ---------------------------
        variant_type_counts = Counter()
        disease_resolution_counts = Counter()
        variant_resolution_counts = Counter()
        pair_status_counts = Counter()

        legacy_diseases = set()
        legacy_rsids = set()
        resolved_current_diseases = set()
        resolved_current_variants = set()

        mapped_rows = []
        unresolved_disease_rows = []
        unresolved_variant_rows = []
        ambiguous_variant_rows = []
        pair_not_in_current_dva = []

        for row_no, row in enumerate(
            ws.iter_rows(min_row=2, values_only=True), start=2
        ):
            did_raw = normalize_oddb_id(cell(row, did_i))
            dname = cell(row, dname_i)
            snpid = cell(row, snpid_i)
            dbsnp_url = cell(row, dbsnp_i)
            rsid = extract_rsid(snpid, dbsnp_url)
            vtype = cell(row, type_i)

            legacy_diseases.add((did_raw or "", dname))
            if rsid:
                legacy_rsids.add(rsid)
            if vtype:
                variant_type_counts[vtype] += 1

            # Resolve current disease exactly.
            disease_hits = set()
            disease_methods = []

            if did_raw in current_ids:
                disease_hits.add(did_raw)
                disease_methods.append("current DiseaseID")

            if did_raw in old_to_current:
                disease_hits.add(old_to_current[did_raw])
                disease_methods.append("retired/source DiseaseID")

            if dname:
                hits = by_name.get(norm_name(dname), set())
                if hits:
                    disease_hits.update(hits)
                    disease_methods.append("exact disease name/alias")

            if len(disease_hits) != 1:
                unresolved_disease_rows.append((
                    row_no, did_raw or "", dname, rsid or "",
                    "|".join(sorted(disease_hits))
                ))
                disease_resolution_counts[
                    "ambiguous" if len(disease_hits) > 1 else "unresolved"
                ] += 1
                continue

            current_did = next(iter(disease_hits))
            resolved_current_diseases.add(current_did)
            disease_resolution_counts["resolved"] += 1

            # Resolve canonical variant by exact rsID.
            if not rsid:
                unresolved_variant_rows.append((
                    row_no, did_raw or "", dname, snpid, "no rsID extracted"
                ))
                variant_resolution_counts["no_rsid"] += 1
                continue

            v_hits = rsid_to_variant.get(rsid, set())

            if not v_hits:
                unresolved_variant_rows.append((
                    row_no, did_raw or "", dname, rsid, "rsID not found in canonical variant"
                ))
                variant_resolution_counts["unresolved"] += 1
                continue

            if len(v_hits) > 1:
                ambiguous_variant_rows.append((
                    row_no, did_raw or "", dname, rsid,
                    "|".join(sorted(v_hits))
                ))
                variant_resolution_counts["ambiguous"] += 1
                continue

            current_vid = next(iter(v_hits))
            resolved_current_variants.add(current_vid)
            variant_resolution_counts["resolved"] += 1

            pair_present = (current_did, current_vid) in current_dva_pairs
            pair_status_counts[
                "current_DVA_match" if pair_present else "not_in_current_DVA"
            ] += 1

            record = {
                "Source_Row": row_no,
                "Legacy_DiseaseID": did_raw or "",
                "Legacy_Disease_Name": dname,
                "Current_DiseaseID": current_did,
                "Disease_Mapping_Method": "+".join(sorted(set(disease_methods))),
                "rsID": rsid,
                "Current_VariantID": current_vid,
                "Current_DVA_Pair": "yes" if pair_present else "no",
                "Disease_Category": cell(row, category_i),
                "Chromosome": cell(row, chr_i),
                "Position_hg38": cell(row, pos_i),
                "Major_Allele": cell(row, allele_i),
                "Variant_Detail": cell(row, detail_i),
                "Variant_Type": vtype,
                "Gene": cell(row, gene_i),
                "Ensembl_URL": cell(row, ens_i),
                "dbSNP_URL": dbsnp_url,
                "gnomAD_URL": cell(row, gnomad_i),
                "DOIs": cell(row, doi_i),
            }
            mapped_rows.append(record)

            if not pair_present:
                pair_not_in_current_dva.append(record)

        observed_types = list(variant_type_counts.keys())
        missing_expected_types = [
            x for x in EXPECTED_VARIANT_TYPES
            if x not in variant_type_counts
        ]
        unexpected_types = [
            x for x in observed_types
            if x not in EXPECTED_VARIANT_TYPES
        ]

        print("ODDB LEGACY VARIANT → CURRENT ODDB v1 MAPPING AUDIT v0.5.1")
        print("=" * 92)
        print(f"Workbook rows: {ws.max_row - 1:,}")
        print(f"Legacy unique disease keys: {len(legacy_diseases):,}")
        print(f"Legacy unique rsIDs: {len(legacy_rsids):,}")
        print(f"Observed variant.type categories: {len(variant_type_counts):,}")
        print(f"Current canonical diseases represented after mapping: "
              f"{len(resolved_current_diseases):,}")
        print(f"Current canonical variants represented after mapping: "
              f"{len(resolved_current_variants):,}")
        print(f"Mapped legacy rows: {len(mapped_rows):,}")

        print("\nDisease resolution:")
        for k, n in disease_resolution_counts.items():
            print(f"{k:20s} {n:7d}")

        print("\nVariant resolution:")
        for k, n in variant_resolution_counts.items():
            print(f"{k:20s} {n:7d}")

        print("\nDisease–variant pair verification:")
        for k, n in pair_status_counts.items():
            print(f"{k:24s} {n:7d}")

        print("\nvariant.type categories:")
        for value, n in variant_type_counts.most_common():
            print(f"{n:7d}  {value}")

        print("\nExpected taxonomy check:")
        print(f"Missing expected types: {missing_expected_types or 'none'}")
        print(f"Unexpected types: {unexpected_types or 'none'}")

        fieldnames = [
            "Source_Row",
            "Legacy_DiseaseID", "Legacy_Disease_Name",
            "Current_DiseaseID", "Disease_Mapping_Method",
            "rsID", "Current_VariantID", "Current_DVA_Pair",
            "Disease_Category", "Chromosome", "Position_hg38",
            "Major_Allele", "Variant_Detail", "Variant_Type", "Gene",
            "Ensembl_URL", "dbSNP_URL", "gnomAD_URL", "DOIs",
        ]

        with (report_dir / "mapped_legacy_variant_rows.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.DictWriter(f, fieldnames=fieldnames, delimiter="\t")
            w.writeheader()
            w.writerows(mapped_rows)

        with (report_dir / "unresolved_disease_rows.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow([
                "Source_Row", "Legacy_DiseaseID", "Legacy_Disease_Name",
                "rsID", "Candidate_Current_DiseaseIDs"
            ])
            w.writerows(unresolved_disease_rows)

        with (report_dir / "unresolved_variant_rows.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow([
                "Source_Row", "Legacy_DiseaseID", "Legacy_Disease_Name",
                "Variant", "Reason"
            ])
            w.writerows(unresolved_variant_rows)

        with (report_dir / "ambiguous_variant_rows.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow([
                "Source_Row", "Legacy_DiseaseID", "Legacy_Disease_Name",
                "rsID", "Candidate_Current_VariantIDs"
            ])
            w.writerows(ambiguous_variant_rows)

        with (report_dir / "mapped_rows_not_in_current_dva.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.DictWriter(f, fieldnames=fieldnames, delimiter="\t")
            w.writeheader()
            w.writerows(pair_not_in_current_dva)

        with (report_dir / "variant_type_counts.tsv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            w = csv.writer(f, delimiter="\t")
            w.writerow(["Variant_Type", "Legacy_Row_Count"])
            for value, n in variant_type_counts.most_common():
                w.writerow([value, n])

        print("\nReports:")
        for name in [
            "mapped_legacy_variant_rows.tsv",
            "unresolved_disease_rows.tsv",
            "unresolved_variant_rows.tsv",
            "ambiguous_variant_rows.tsv",
            "mapped_rows_not_in_current_dva.tsv",
            "variant_type_counts.tsv",
        ]:
            print(report_dir / name)

        print("\nREAD-ONLY: MySQL and the legacy workbook were not modified.")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
